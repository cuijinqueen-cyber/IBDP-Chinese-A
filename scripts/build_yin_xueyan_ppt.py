#!/usr/bin/env python3
"""Generate a lecture PPT on literary concepts and techniques in Bai Xianyong's
《永远的尹雪艳》, organized from text surface to theme."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# --- canvas ---
W = Inches(13.333)
H = Inches(7.5)

# --- palette: ivory paper, ink, snow-white, blood-red ---
IVORY = RGBColor(0xF6, 0xF0, 0xE4)
INK = RGBColor(0x1E, 0x18, 0x16)
INK2 = RGBColor(0x32, 0x2A, 0x26)
WINE = RGBColor(0x4A, 0x16, 0x1C)
CRIMSON = RGBColor(0xA3, 0x24, 0x32)
SOFT_RED = RGBColor(0xC4, 0x5A, 0x5E)
GOLD = RGBColor(0xB8, 0x93, 0x5A)
MUTED = RGBColor(0x6F, 0x65, 0x5C)
WARM_GRAY = RGBColor(0x8A, 0x80, 0x76)
CARD = RGBColor(0xFF, 0xFB, 0xF5)
SNOW = RGBColor(0xEE, 0xE8, 0xDE)
LINE = RGBColor(0xD9, 0xCF, 0xC0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# PowerPoint-safe CJK font (falls back on Mac/Linux)
FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
blank = prs.slide_layouts[6]


def set_run(run, text, size, color, bold=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn("a:%s" % tag))
        if el is None:
            el = etree.SubElement(rPr, qn("a:%s" % tag))
        el.set("typeface", FONT)


def add_rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def add_round(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.adjustments[0] = 0.08
    sh.shadow.inherit = False
    return sh


def tb(slide, l, t, w, h, text, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    set_run(r, text, size, color, bold)
    return box


def tb_multi(slide, l, t, w, h, lines, default_size=15, default_color=INK, align=PP_ALIGN.LEFT, spacing=1.05):
    """lines: list of str or dict {text, size, color, bold}"""
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        p.line_spacing = spacing
        if isinstance(item, str):
            r = p.add_run()
            set_run(r, item, default_size, default_color, False)
        else:
            r = p.add_run()
            set_run(
                r,
                item.get("text", ""),
                item.get("size", default_size),
                item.get("color", default_color),
                item.get("bold", False),
            )
    return box


def bg(slide, color=IVORY):
    add_rect(slide, 0, 0, W, H, color)


def footer(slide, page, total=33, light=False):
    c = RGBColor(0xC8, 0xBE, 0xB2) if not light else RGBColor(0x8A, 0x70, 0x72)
    tb(slide, Inches(0.5), Inches(7.18), Inches(9), Inches(0.28),
       "《永远的尹雪艳》· 文学概念与手法  |  白先勇《台北人》", 10, c)
    tb(slide, Inches(11.4), Inches(7.18), Inches(1.5), Inches(0.28),
       f"{page}  /  {total}", 10, c, align=PP_ALIGN.RIGHT)


def section_bar(slide, kicker, title, subtitle=None):
    add_rect(slide, 0, 0, W, Inches(0.08), CRIMSON)
    tb(slide, Inches(0.55), Inches(0.22), Inches(12), Inches(0.28), kicker, 12, CRIMSON, True)
    tb(slide, Inches(0.55), Inches(0.46), Inches(12), Inches(0.48), title, 28, INK, True)
    if subtitle:
        tb(slide, Inches(0.55), Inches(0.94), Inches(12), Inches(0.32), subtitle, 13, MUTED)


def card(slide, l, t, w, h, title, body_lines, accent=CRIMSON, title_size=15, body_size=12):
    add_round(slide, l, t, w, h, CARD, LINE)
    add_rect(slide, l, t, Inches(0.08), h, accent)
    tb(slide, l + Inches(0.22), t + Inches(0.12), w - Inches(0.35), Inches(0.36),
       title, title_size, INK, True)
    tb_multi(
        slide,
        l + Inches(0.22),
        t + Inches(0.48),
        w - Inches(0.38),
        h - Inches(0.58),
        body_lines,
        default_size=body_size,
        default_color=INK2,
    )


# ============================================================
# SLIDE 1  Cover
# ============================================================
s = prs.slides.add_slide(blank)
bg(s, WINE)
add_rect(s, 0, 0, Inches(0.18), H, GOLD)
add_rect(s, 0, Inches(6.95), W, Inches(0.55), RGBColor(0x36, 0x10, 0x14))
tb(s, Inches(0.7), Inches(1.15), Inches(12), Inches(0.35),
   "IBDP Chinese A  ·  文学复习讲座", 14, GOLD, True)
tb(s, Inches(0.7), Inches(1.7), Inches(12), Inches(1.1),
   "永远的尹雪艳", 48, WHITE, True)
tb(s, Inches(0.7), Inches(2.85), Inches(12), Inches(0.45),
   "文学概念与手法梳理", 26, RGBColor(0xF0, 0xD4, 0xC0))
tb(s, Inches(0.7), Inches(3.5), Inches(11.5), Inches(0.7),
   "从人物形象出发，经塑造、叙事、空白、意象与语言，落到主旨。\n按「文本表层 → 艺术手法 → 象征意义 → 主题思想」的逻辑展开。",
   15, RGBColor(0xE8, 0xC8, 0xC4))
tb(s, Inches(0.7), Inches(5.55), Inches(8), Inches(0.7),
   "白先勇  ·  《台北人》首篇\n短篇小说  ·  1960年代台北 / 旧上海记忆", 14, RGBColor(0xD4, 0xB0, 0xA8))
tb(s, Inches(0.7), Inches(7.05), Inches(12), Inches(0.3),
   "概念界定  ·  手法识别  ·  例证落点  ·  主旨贯通", 12, GOLD)

# ============================================================
# SLIDE 2  为何这样排
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "阅读路径", "为什么按这个顺序讲？", "文学分析不是清单堆砌，而是从「看见什么」走向「如何做成」再走向「意味着什么」。")
steps = [
    ("01", "定位文本", "文类、系列、核心句子\n「尹雪艳总也不老」"),
    ("02", "抓住人物", "形象是小说的枢纽\n矛盾体 → 象征体"),
    ("03", "拆解手法", "塑造 · 叙事 · 空白\n意象 · 语言 · 反讽"),
    ("04", "落到主旨", "时间、欲望、命运\n怀旧与虚无"),
]
for i, (n, t, b) in enumerate(steps):
    x = Inches(0.5 + i * 3.2)
    add_round(s, x, Inches(1.55), Inches(3.0), Inches(3.55), CARD, LINE)
    tb(s, x + Inches(0.22), Inches(1.75), Inches(2.5), Inches(0.45), n, 22, CRIMSON, True)
    tb(s, x + Inches(0.22), Inches(2.25), Inches(2.55), Inches(0.4), t, 18, INK, True)
    tb(s, x + Inches(0.22), Inches(2.8), Inches(2.55), Inches(1.9), b, 14, MUTED)
    if i < 3:
        tb(s, x + Inches(2.85), Inches(3.0), Inches(0.4), Inches(0.4), "→", 20, GOLD, True)
tb(s, Inches(0.55), Inches(5.35), Inches(12.2), Inches(1.5),
   "贯穿原则：反讽。它既是人物手法，也是叙事策略，更是结构原则——表面赞美、内里批判。\n"
   "因此本讲不把「反讽」提前单列成第一章，而让它在塑造、叙事、语言中反复现身，最后收束为整体结构。",
   14, INK2)
footer(s, 2)

# ============================================================
# SLIDE 3  目录
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "CONTENTS", "八个板块")
items = [
    ("一", "作品定位与核心问题", "文类 · 《台北人》 · 开篇句"),
    ("二", "人物形象：矛盾与象征", "雪/艳 · 冷热 · 神/妖/死神"),
    ("三", "人物塑造手法", "以形写神 · 白描 · 语言 · 细节"),
    ("四", "叙事手法", "插叙 · 全知 · 不可靠叙述"),
    ("五", "空白手法", "伊瑟尔 · 人物/环境/情节空白"),
    ("六", "环境、意象与象征", "白与红 · 晚香玉 · 麻将 · 公馆"),
    ("七", "语言特色与反讽结构", "冷静 · 暗示 · 比喻 · 整体反讽"),
    ("八", "主旨与答题路径", "六层主旨 · 手法如何服务主题"),
]
for i, (n, t, b) in enumerate(items):
    col = i % 2
    row = i // 2
    x = Inches(0.5 + col * 6.4)
    y = Inches(1.5 + row * 1.28)
    add_round(s, x, y, Inches(6.15), Inches(1.12), CARD, LINE)
    add_rect(s, x, y, Inches(0.1), Inches(1.12), CRIMSON if i % 2 == 0 else GOLD)
    tb(s, x + Inches(0.28), y + Inches(0.16), Inches(0.6), Inches(0.4), n, 20, CRIMSON, True)
    tb(s, x + Inches(0.95), y + Inches(0.18), Inches(4.9), Inches(0.4), t, 16, INK, True)
    tb(s, x + Inches(0.95), y + Inches(0.58), Inches(4.9), Inches(0.38), b, 12, MUTED)
footer(s, 3)

# ============================================================
# SLIDE 4  作品定位
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "01  作品定位", "先把文本放进正确的坐标")
cards4 = [
    ("文类", ["短篇小说，情节完整、有头有尾", "传统线性骨架 + 插叙补层", "适合考查：视角、象征、反讽"]),
    ("系列位置", ["《台北人》十四篇之首", "为整部集子定下调性：怀旧、沧桑、嘲讽", "「永远」与「台北人」的时间张力"]),
    ("时空", ["现在：1960年代台北尹公馆", "过去：上海百乐门、霞飞路", "视点立足现在，不断插入上海"]),
    ("核心句子", ["开篇：「尹雪艳总也不老。」", "既是人物设定，也是反讽基调", "谁能不老？「永远」是幻觉"]),
]
for i, (t, lines) in enumerate(cards4):
    x = Inches(0.45 + (i % 4) * 3.2)
    y = Inches(1.5)
    add_round(s, x, y, Inches(3.05), Inches(3.7), CARD, LINE)
    add_rect(s, x, y, Inches(3.05), Inches(0.08), CRIMSON if i != 3 else GOLD)
    tb(s, x + Inches(0.2), y + Inches(0.25), Inches(2.65), Inches(0.4), t, 16, CRIMSON, True)
    tb_multi(s, x + Inches(0.2), y + Inches(0.8), Inches(2.65), Inches(2.6),
             [{"text": "·  " + ln, "size": 13} for ln in lines], default_size=13, default_color=INK2)
tb(s, Inches(0.55), Inches(5.45), Inches(12.2), Inches(1.4),
   "分析入口：不要一上来赞美「永远的美人」。先问——叙述者为什么这样说？隐含作者是否同意？\n"
   "「总也不老」是表层话语；「谁能不老」才是文本真正要读者想到的常识。二者的裂隙，就是反讽开始的地方。",
   14, INK2)
footer(s, 4)

# ============================================================
# SLIDE 5  概念地图
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "01  作品定位", "概念地图：手法如何通向主旨")
# flow boxes
layers = [
    (Inches(1.45), "人物形象", "尹雪艳：雪/艳 · 冷/热 · 人/神/死神"),
    (Inches(2.45), "塑造手法", "以形写神 · 白描 · 语言 · 细节"),
    (Inches(3.45), "叙事与空白", "插叙 · 全知冷静 · 不可靠叙述 · 留白"),
    (Inches(4.45), "意象与语言", "白/红 · 晚香玉 · 麻将 · 客观冷静 · 暗示"),
    (Inches(5.45), "反讽结构", "话语反讽 · 情景反讽 · 整体反讽"),
]
for y, title, body in layers:
    add_round(s, Inches(0.7), y, Inches(8.3), Inches(0.85), CARD, LINE)
    tb(s, Inches(0.95), y + Inches(0.08), Inches(2.4), Inches(0.65), title, 15, CRIMSON, True, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, Inches(3.4), y + Inches(0.08), Inches(5.4), Inches(0.65), body, 13, INK2, anchor=MSO_ANCHOR.MIDDLE)

add_round(s, Inches(9.3), Inches(1.45), Inches(3.5), Inches(4.85), WINE)
tb(s, Inches(9.5), Inches(1.7), Inches(3.1), Inches(0.4), "落点：主旨", 16, GOLD, True)
aims = ["时间与命运不可抗拒", "繁华与衰落的对照", "人性的复杂与矛盾", "怀旧与现实的冲突", "女性的力量与悲剧", "欲望、轮回与虚无"]
tb_multi(s, Inches(9.5), Inches(2.25), Inches(3.1), Inches(3.7),
         [{"text": "▸  " + a, "size": 13, "color": WHITE} for a in aims], default_color=WHITE)
footer(s, 5)

# ============================================================
# SLIDE 6  人物：为何是枢纽
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "02  人物形象", "尹雪艳是整篇小说的概念枢纽", "先分清三层：作为「人」的性格、作为「符号」的象征、作为「叙述效果」的神秘。")
layers3 = [
    ("表层：可见的人", CRIMSON,
     ["容颜不老、银白旗袍、风情", "社交女王：牌局、公馆、应酬", "淡然世故、八面玲珑", "对男女皆有吸引力"]),
    ("中层：矛盾的性格", GOLD,
     ["冷艳 ↔ 慈悲（安抚失意者）", "超然 ↔ 寄生（吃红、周旋）", "圣洁（雪）↔ 魅惑（艳）", "女祭司 / 命运女神 / 红颜祸水"]),
    ("深层：象征的功能", INK,
     ["时间：永恒 vs 变迁", "欲望：心魔与黑洞", "命运：摆渡者 / 死神", "旧上海：海市蜃楼式繁华"]),
]
for i, (t, acc, lines) in enumerate(layers3):
    x = Inches(0.5 + i * 4.2)
    add_round(s, x, Inches(1.55), Inches(4.0), Inches(4.55), CARD, LINE)
    add_rect(s, x, Inches(1.55), Inches(4.0), Inches(0.1), acc)
    tb(s, x + Inches(0.25), Inches(1.8), Inches(3.5), Inches(0.45), t, 17, acc if acc != INK else CRIMSON, True)
    tb_multi(s, x + Inches(0.25), Inches(2.4), Inches(3.5), Inches(3.4),
             [{"text": "·  " + ln, "size": 14} for ln in lines], default_size=14)
footer(s, 6)

# ============================================================
# SLIDE 7  雪与艳
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "02  人物形象", "名字即概念：「雪」与「艳」的矛盾统一")
add_round(s, Inches(0.5), Inches(1.5), Inches(5.9), Inches(4.7), CARD, LINE)
tb(s, Inches(0.8), Inches(1.7), Inches(5.3), Inches(0.4), "雪  ·  冷、白、圣洁、死亡", 18, RGBColor(0x5A, 0x6A, 0x78), True)
snow_pts = [
    "雪白肌肤、素白旗袍、银狐大氅",
    "「冰雪化成的精灵」「通身银白的女祭司」",
    "不擦胭抹粉、不穿红戴绿",
    "内心冷酷、超然、均衡不受外界迁异",
    "白色 = 死亡寒意、花圈丧幛",
]
tb_multi(s, Inches(0.8), Inches(2.3), Inches(5.3), Inches(3.5),
         [{"text": "·  " + p, "size": 15} for p in snow_pts])

add_round(s, Inches(6.7), Inches(1.5), Inches(6.1), Inches(4.7), CARD, LINE)
tb(s, Inches(7.0), Inches(1.7), Inches(5.5), Inches(0.4), "艳  ·  红、热、魅惑、血光", 18, CRIMSON, True)
yan_pts = [
    "血红郁金香、鲜红樱桃、桃花心红木",
    "「世人不及的风情」、舞池中的从容",
    "重煞传说：轻则家败，重则人亡",
    "吸引王贵生、洪处长、徐壮图前仆后继",
    "红色 = 欲望、血腥、见血预兆",
]
tb_multi(s, Inches(7.0), Inches(2.3), Inches(5.5), Inches(3.5),
         [{"text": "·  " + p, "size": 15} for p in yan_pts])
footer(s, 7)

# ============================================================
# SLIDE 8  神妖死神
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "02  人物形象", "介于神灵与妖孽之间，导向「死神」")
triples = [
    ("神灵一侧", "通身银白的女祭司\n命运女神般安抚客人们\n悲天悯人地看牌桌厮杀\n对吴经理、宋太太的慈悲", GOLD),
    ("妖孽一侧", "八字带重煞、犯白虎\n吴家阿婆比作妲己、褒姒\n「狐狸精」「祸水」的流言\n所到之处「血光一片」", CRIMSON),
    ("死神化身", "不老不死，旁观众生老去\n「我来吃你的红」的双关\n加速纵欲者走向毁灭\n命运舟楫上的清冷摆渡者", INK),
]
for i, (t, b, acc) in enumerate(triples):
    x = Inches(0.5 + i * 4.2)
    add_round(s, x, Inches(1.5), Inches(4.0), Inches(3.7), CARD, LINE)
    add_rect(s, x, Inches(1.5), Inches(4.0), Inches(0.08), acc)
    tb(s, x + Inches(0.25), Inches(1.75), Inches(3.5), Inches(0.4), t, 17, acc if acc != INK else CRIMSON, True)
    tb(s, x + Inches(0.25), Inches(2.3), Inches(3.5), Inches(2.6), b, 14, INK2)
tb(s, Inches(0.55), Inches(5.45), Inches(12.2), Inches(1.4),
   "关键判断：她不是「坏女人」这么简单。神、妖、死神三种读法同时成立，人物才有张力。\n"
   "她的永恒 vs 周围人事的无常、她的清冷 vs 众人的炙热追逐——变与不变反复对照，指向欲望腐蚀与命运无助。",
   14, INK2)
footer(s, 8)

# ============================================================
# SLIDE 9  社交掌控与双重
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "02  人物形象", "社交掌控力，与冷艳/慈悲的双重性格")
card(s, Inches(0.5), Inches(1.5), Inches(6.1), Inches(5.0), "掌控：公馆即舞台",
     [
         {"text": "·  百乐门到尹公馆，她始终是全场中心", "size": 14},
         {"text": "·  懂牌品癖性，巧妙排局，不伤和气", "size": 14},
         {"text": "·  用「作废了的头衔」称呼旧客，恢复其优越感", "size": 14},
         {"text": "·  对徐壮图：劝酒、点牌、特别款待", "size": 14},
         {"text": "·  见什么人说什么话——金口难开，开口即中", "size": 14},
         {"text": "", "size": 8},
         {"text": "概念：人物的「功能性」——她是社交机器，也是怀旧药剂。", "size": 13, "color": CRIMSON, "bold": True},
     ], title_size=16, body_size=14)
card(s, Inches(6.85), Inches(1.5), Inches(6.0), Inches(5.0), "双重：冷艳与慈悲",
     [
         {"text": "冷：素白旗袍、距离感的微笑、处理后事的淡漠", "size": 14},
         {"text": "慈：电暖炉与铁观音给吴经理；热毛巾给宋太太", "size": 14},
         {"text": "「人无千日好，花无百日红」——像神谕，也像麻醉", "size": 14},
         {"text": "", "size": 8},
         {"text": "读法提醒：慈悲可以是真的同情，也可以是经营人头、维持牌局的手腕。文本允许两种理解并存——这正是空白与反讽的接口。", "size": 13, "color": MUTED},
     ], GOLD, title_size=16, body_size=14)
footer(s, 9)

# ============================================================
# SLIDE 10  象征网络
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "02  人物形象", "尹雪艳的象征网络（答题时按题干选一层，勿全堆）")
syms = [
    ("时间的旁观者", "自己恒定，见证他人老去。她是刻度，标出过去与现在的界限。"),
    ("旧时代的残影", "公馆 = 上海繁华的仿制品；她周身「透着荣华的麝香」。"),
    ("欲望的外化 / 心魔", "男人不是被煞气所害，是被自己的欲望吞没。她是导火索与镜子。"),
    ("命运 / 死神", "摆渡者、女祭司、吃红。纵欲加速死亡；她不悲不喜。"),
    ("人性的清醒与冷酷", "旁观厮杀，自己不入局——理智，也是非人。"),
    ("乌托邦与虚无", "尹公馆是精神欲望的乌托邦；繁华背后是枯骨，色即是空。"),
]
for i, (t, b) in enumerate(syms):
    col, row = i % 3, i // 3
    x = Inches(0.45 + col * 4.25)
    y = Inches(1.5 + row * 2.5)
    add_round(s, x, y, Inches(4.05), Inches(2.28), CARD, LINE)
    tb(s, x + Inches(0.22), y + Inches(0.2), Inches(3.6), Inches(0.45), f"{i+1}  {t}", 15, CRIMSON, True)
    tb(s, x + Inches(0.22), y + Inches(0.75), Inches(3.6), Inches(1.3), b, 13, INK2)
footer(s, 10)

# ============================================================
# SLIDE 11  塑造手法总览
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "03  人物塑造手法", "四种核心手法：由外而内，由形到神")
techs = [
    ("1", "以形写神", "衣着、神情、动作\n写出内心的「冷」与沧桑"),
    ("2", "白描勾勒", "抓住突出特征\n三言两语，神情毕现"),
    ("3", "个性化语言", "见人说话、金口难开\n话里露出寄生本质"),
    ("4", "特征性细节", "郁金香颤动、抚摩孩子\n当晚成牌局 → 心理与个性"),
]
for i, (n, t, b) in enumerate(techs):
    x = Inches(0.5 + i * 3.2)
    add_round(s, x, Inches(1.55), Inches(3.05), Inches(3.5), CARD, LINE)
    tb(s, x + Inches(0.22), Inches(1.75), Inches(2.6), Inches(0.45), n, 22, CRIMSON, True)
    tb(s, x + Inches(0.22), Inches(2.25), Inches(2.6), Inches(0.5), t, 17, INK, True)
    tb(s, x + Inches(0.22), Inches(2.9), Inches(2.6), Inches(1.8), b, 14, MUTED)
tb(s, Inches(0.55), Inches(5.3), Inches(12.2), Inches(1.5),
   "概念辨析：「以形写神」是中国传统小说/画论术语，强调外形描写通向精神气质；「白描」强调少形容词、抓特征。\n"
   "二者常叠用：开篇「总也不老」+ 雪白肌肤、细挑身材，既是白描，也在以形写神。答题时点明侧重点即可。",
   14, INK2)
footer(s, 11)

# ============================================================
# SLIDE 12  以形写神
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "03  人物塑造手法", "以形写神：衣着之「冷」与神情之「悯」")
card(s, Inches(0.5), Inches(1.5), Inches(6.1), Inches(5.0), "形：反复写「银白」",
     [
         {"text": "概念：通过服饰、动作的精细描绘，让性格从外形中长出来。", "size": 13, "color": MUTED},
         {"text": "", "size": 6},
         {"text": "·  夏天浑身银白，净扮了不得", "size": 14},
         {"text": "·  盛宴：翻领束腰银狐大氅", "size": 14},
         {"text": "·  寿宴：月白织锦旗袍", "size": 14},
         {"text": "·  徐壮图祭悼：仍一身素白", "size": 14},
         {"text": "", "size": 8},
         {"text": "一个字：冷。衣着之冷衬容貌之艳，更暗示内心冷酷；情节推进后，外冷与内冷完全融合。", "size": 13, "color": CRIMSON, "bold": True},
     ])
card(s, Inches(6.85), Inches(1.5), Inches(6.0), Inches(5.0), "神：喷烟圈与悲天悯人",
     [
         {"text": "麻将桌上，客人们互相厮杀时：", "size": 13, "color": MUTED},
         {"text": "", "size": 6},
         {"text": "「叼着金嘴子的三个九，徐徐地喷着烟圈，以悲天悯人的眼光看着她这一群……狂热的互相厮杀」", "size": 13},
         {"text": "", "size": 8},
         {"text": "效果：冷静、清醒、局外人姿态；同时「这一群」是她的，她悲悯他们也悲悯自己——美人迟暮的潜台词由此写出。", "size": 13, "color": INK2},
     ], GOLD)
footer(s, 12)

# ============================================================
# SLIDE 13  白描
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "03  人物塑造手法", "白描：洗练勾勒，以少胜多")
tb(s, Inches(0.55), Inches(1.45), Inches(12), Inches(0.45),
   "概念：抓住最突出的特征，用简洁文字勾勒，三言两语神情毕现。不靠堆砌形容词，而靠「选哪几笔」。", 14, MUTED)
card(s, Inches(0.5), Inches(2.0), Inches(6.1), Inches(4.5), "例：吴经理",
     [
         {"text": "头发全白 · 风湿蹒跚 · 沙眼倒睫 · 眼圈溃烂露粉红肉", "size": 14, "bold": True, "color": CRIMSON},
         {"text": "", "size": 8},
         {"text": "几笔下去：暮路英雄、晚景凄凉。读者自然对照他「银行总经理」的过去，人世沧桑不必作者直说。", "size": 14},
         {"text": "", "size": 8},
         {"text": "功能：用「老」衬托尹雪艳的「不老」；为后文「老当益壮」的反讽预埋事实压力。", "size": 14},
     ])
card(s, Inches(6.85), Inches(2.0), Inches(6.0), Inches(4.5), "例：尹雪艳开篇",
     [
         {"text": "「尹雪艳总也不老。」", "size": 14, "bold": True, "color": CRIMSON},
         {"text": "雪白肌肤、细挑身材、俏丽恬静眉眼；「像一株晚开的玉梨花」。", "size": 14},
         {"text": "", "size": 8},
         {"text": "精练描写突出美丽，也暗示冷酷（玉、晚开、压倒群芳）。", "size": 14},
         {"text": "白描在这里同时完成「美」的建立与「非人」的预感。", "size": 14},
     ], GOLD)
footer(s, 13)

# ============================================================
# SLIDE 14  语言写性格
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "03  人物塑造手法", "通过人物语言写性格：多声部，而非定音定调")
rows = [
    ("对吴经理", "「干爹才是老当益壮呢！」", "安慰 + 经营：他能引来商界人头。话语反讽：与风烛残年的外形完全相反。"),
    ("对旧客", "用十几年前作废的头衔称呼", "心理按摩：优越感「如同受过诰封」。她贩卖的是身份幻觉。"),
    ("对太太们", "「人无千日好，花无百日红」", "安抚焦躁，也把她们留在公馆里「烘云托月」。"),
    ("对徐壮图", "「稀客……自然要跟别人不同一点」", "高雅熨帖，交际手腕；特别款待即诱惑的开始。"),
    ("收束句", "「回头赢了……我来吃你的红。」", "双关：吃彩头 / 死神吞没。寄生本质与死亡主题叠合。"),
]
add_round(s, Inches(0.45), Inches(1.45), Inches(12.4), Inches(5.2), CARD, LINE)
headers = [("对象", 0.2), ("原话 / 做法", 2.0), ("性格与功能", 6.6)]
for t, xoff in headers:
    tb(s, Inches(0.65 + xoff), Inches(1.55), Inches(4 if xoff else 1.7), Inches(0.35), t, 12, CRIMSON, True)
add_rect(s, Inches(0.65), Inches(1.95), Inches(12.0), Inches(0.015), LINE)
for i, (a, b, c) in enumerate(rows):
    y = Inches(2.1 + i * 0.85)
    tb(s, Inches(0.65), y, Inches(1.7), Inches(0.75), a, 13, INK, True)
    tb(s, Inches(2.45), y, Inches(4.4), Inches(0.75), b, 13, CRIMSON)
    tb(s, Inches(7.05), y, Inches(5.5), Inches(0.75), c, 13, INK2)
footer(s, 14)

# ============================================================
# SLIDE 15  细节
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "03  人物塑造手法", "特征性细节：一小处动作，托出心理与个性")
dets = [
    ("血红郁金香「颤巍巍地抖动」", "初识徐壮图。细节把内心的热写到头饰上——白衣上的红，欲望破例外露。"),
    ("银耳坠「来回地浪荡」", "从徐壮图眼中写出。诱惑被对象化，读者与徐壮图同时「着道」。"),
    ("葬礼：抚摩孩子的头，与徐太太握手", "庄重得体的社交表演；「良心」的仪节，而非哀伤。"),
    ("当晚又成上牌局，有人是祭悼会后约好的", "虚伪、狠冷、循环开始。新的「徐壮图」已经在路上。"),
]
for i, (t, b) in enumerate(dets):
    y = Inches(1.48 + i * 1.25)
    add_round(s, Inches(0.5), y, Inches(12.3), Inches(1.12), CARD, LINE)
    add_rect(s, Inches(0.5), y, Inches(0.1), Inches(1.12), CRIMSON if i % 2 == 0 else GOLD)
    tb(s, Inches(0.85), y + Inches(0.12), Inches(11.7), Inches(0.35), t, 15, INK, True)
    tb(s, Inches(0.85), y + Inches(0.5), Inches(11.7), Inches(0.5), b, 13, MUTED)
footer(s, 15)

# ============================================================
# SLIDE 16  叙事总览
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "04  叙事手法", "三条叙事线索：结构、视角、可靠性")
nar = [
    ("插叙", "结构", "总体按时间顺序：来台北 → 结识徐壮图 → 丧生 → 公馆依旧繁荣。\n视点立足「现在」，不断插入上海往事，上海与台北、过去与现在交织。"),
    ("全知叙事", "视角", "冷眼旁观，不探入任一角色意识，只写外貌言行与情节。\n刻意不同于白先勇惯用的内心透视，以制造尹雪艳的「冰冷神秘」。"),
    ("不可靠叙述", "叙述者", "叙述者口吻倾慕、赞美；隐含作者冷静批判。\n表层「永远迷人」与深层「不祥、自欺」构成韦恩·布斯所谓的不可信叙述。"),
]
for i, (t, tag, b) in enumerate(nar):
    y = Inches(1.5 + i * 1.65)
    add_round(s, Inches(0.5), y, Inches(12.3), Inches(1.5), CARD, LINE)
    tb(s, Inches(0.8), y + Inches(0.18), Inches(2.4), Inches(0.4), t, 18, CRIMSON, True)
    tb(s, Inches(3.3), y + Inches(0.22), Inches(1.5), Inches(0.35), tag, 12, GOLD, True)
    tb(s, Inches(0.8), y + Inches(0.65), Inches(11.7), Inches(0.7), b, 14, INK2)
footer(s, 16)

# ============================================================
# SLIDE 17  插叙
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "04  叙事手法", "插叙：现在与过去、台北和上海的交织")
add_round(s, Inches(0.5), Inches(1.5), Inches(6.1), Inches(4.95), CARD, LINE)
tb(s, Inches(0.75), Inches(1.7), Inches(5.6), Inches(0.4), "主线（现在 · 台北）", 16, CRIMSON, True)
nows = ["尹雪艳到台北后仍像当年那样迷人", "结识新兴实业巨子徐壮图", "徐壮图因迷她而丧生", "死后尹公馆仍像以前一样繁荣"]
tb_multi(s, Inches(0.75), Inches(2.25), Inches(5.6), Inches(3.8),
         [{"text": f"{i+1}.  {x}", "size": 15} for i, x in enumerate(nows)])
add_round(s, Inches(6.85), Inches(1.5), Inches(6.0), Inches(4.95), CARD, LINE)
tb(s, Inches(7.1), Inches(1.7), Inches(5.5), Inches(0.4), "插入层（过去 · 上海）", 16, GOLD, True)
thens = ["百乐门舞皇后、玉梨花压倒群芳", "王贵生官商勾结下狱枪毙", "洪处长休妻弃子，后丢官破产", "尹公馆对霞飞路排场的刻意维持"]
tb_multi(s, Inches(7.1), Inches(2.25), Inches(5.5), Inches(2.6),
         [{"text": f"{i+1}.  {x}", "size": 15} for i, x in enumerate(thens)])
tb(s, Inches(7.1), Inches(5.1), Inches(5.5), Inches(1.1),
   "效果：社会历史变迁进入命运；「永远」被证明是复制与循环，不是真的静止。", 14, CRIMSON)
footer(s, 17)

# ============================================================
# SLIDE 18  全知
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "04  叙事手法", "全知叙事与冷静嘲讽（语调 tone）")
tb(s, Inches(0.55), Inches(1.42), Inches(12.2), Inches(0.7),
   "欧阳子：《永远的尹雪艳》是《台北人》中嘲讽意味最浓的一篇。作者完全隔离，冷眼旁观；嘲讽前后一贯，藉由叙述者口吻传达。",
   14, MUTED)
pts = [
    ("有所保留的全知", "不进入任何人的内心。尹雪艳因此无法被「理解」为普通人，神秘感成立，象征性才站得住。"),
    ("表层话语", "叙述者只写她美丽、称职、永远浅笑。表层含义必须先成立，隐含作者的批判才能与之对立。"),
    ("语调即反讽通道", "客观冷静 ≠ 没有态度。越是不动声色地记录「宾至如归」「心甘情愿」，谴责与同情越沉郁。"),
    ("与白先勇习惯对照", "他通常展开内心以博取认同与悲悯；本篇反其道，正是为「冰冷」与「嘲讽」服务。"),
]
for i, (t, b) in enumerate(pts):
    col, row = i % 2, i // 2
    x = Inches(0.5 + col * 6.4)
    y = Inches(2.2 + row * 2.15)
    add_round(s, x, y, Inches(6.15), Inches(2.0), CARD, LINE)
    tb(s, x + Inches(0.25), y + Inches(0.2), Inches(5.7), Inches(0.4), t, 16, CRIMSON, True)
    tb(s, x + Inches(0.25), y + Inches(0.7), Inches(5.7), Inches(1.1), b, 14, INK2)
footer(s, 18)

# ============================================================
# SLIDE 19  不可靠叙述
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "04  叙事手法", "不可靠叙述：叙述者 ≠ 隐含作者")
add_round(s, Inches(0.5), Inches(1.45), Inches(12.3), Inches(1.15), WINE)
tb(s, Inches(0.75), Inches(1.55), Inches(11.8), Inches(0.95),
   "韦恩·布斯《小说修辞学》：叙述者若为作品的思想规范（隐含作者的规范）辩护或接近它，便是可靠的；反之不可靠。\n"
   "本篇叙述者站在「着了道的人」一边赞美她；隐含作者则冷静、批判。裂隙即反讽。",
   14, WHITE)
card(s, Inches(0.5), Inches(2.8), Inches(6.1), Inches(3.7), "叙述者看见的",
     [
         {"text": "「不管人事怎么变迁，尹雪艳永远是尹雪艳」", "size": 13, "bold": True},
         {"text": "「迷人的地方实在讲不清，数不尽」", "size": 13},
         {"text": "「最称职的主人」「还算有良心」", "size": 13},
         {"text": "", "size": 6},
         {"text": "这是被蛊惑者的感受。读者若停在这一层，会误读为赞美佳人。", "size": 13, "color": MUTED},
     ])
card(s, Inches(6.85), Inches(2.8), Inches(6.0), Inches(3.7), "隐含作者埋下的「漏洞」",
     [
         {"text": "王贵生枪毙那天，她只在百乐门「停了一宵」", "size": 13},
         {"text": "离洪处长「还算有良心」——带走家当与仆人", "size": 13},
         {"text": "徐壮图葬礼当晚即开牌局", "size": 13},
         {"text": "", "size": 6},
         {"text": "敏感读者从自相矛盾中读出言外之意：薄情、虚伪、循环。", "size": 13, "color": CRIMSON, "bold": True},
     ], GOLD)
footer(s, 19)

# ============================================================
# SLIDE 20  不可靠的目的
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "04  叙事手法", "不可靠叙述的两个目的")
add_round(s, Inches(0.5), Inches(1.5), Inches(6.1), Inches(5.0), CARD, LINE)
tb(s, Inches(0.75), Inches(1.7), Inches(5.6), Inches(0.45), "一、让读者先「着道」", 18, CRIMSON, True)
tb(s, Inches(0.75), Inches(2.3), Inches(5.6), Inches(3.9),
   "尹雪艳的魅力依赖视听嗅觉，文字难以直接传递，故叙述者用夸张比喻：\n\n"
   "「冰雪化成的精灵，冷艳逼人」\n「像一阵三月的微风……全场的人都好像给这阵风熏中了」\n\n"
   "读者须先感到难以抗拒，才能理解悲剧：人难以抗拒她，正如难以抗拒欲望与死亡。这是人类本能弱点造成的永恒悲剧。",
   14, INK2)
add_round(s, Inches(6.85), Inches(1.5), Inches(6.0), Inches(5.0), CARD, LINE)
tb(s, Inches(7.1), Inches(1.7), Inches(5.5), Inches(0.45), "二、使象征性浮现", 18, GOLD, True)
tb(s, Inches(7.1), Inches(2.3), Inches(5.5), Inches(3.9),
   "反讽造成形象的不确定性：「永远不老」「血光一片」都不像真人。\n\n"
   "她象征心魔——永远无法满足却不由自主受其吸引的欲望；过度欲望加速毁灭。\n\n"
   "尹雪艳像镜子，反照出人们不肯正视台北荒凉与身份跌落、自欺欺人的真面目。",
   14, INK2)
footer(s, 20)

# ============================================================
# SLIDE 21  空白概念
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "05  空白手法", "伊瑟尔：文本是充满「未定点」与「空白」的结构")
add_round(s, Inches(0.5), Inches(1.48), Inches(12.3), Inches(1.35), WINE)
tb(s, Inches(0.75), Inches(1.65), Inches(11.8), Inches(1.05),
   "沃尔夫冈·伊瑟尔（接受美学）：作品意义不在文本里封存完毕，而在阅读中对空白的填补。\n"
   "本篇的空白是综合性手法——可以是人物空白、环境空白，也可以是情节空白。它服务于象征，而不是故弄玄虚。",
   15, WHITE)
kinds = [
    ("人物空白", "正面细描极少，年龄长相难以坐实。侧击、传说、留白如水墨，想象空间被故意撑开。开篇把她推向「非人化」。"),
    ("环境空白", "公馆写得很满（红木、黑枕、晚香玉），尹雪艳本人却是素白「留白」。满与空对照，追问她究竟是谁。"),
    ("情节 / 心理空白", "三名男人结局写了，她的态度只用「算是」「还算」「当晚」。全文几乎没有她的心理——女主人公「失声」。"),
]
for i, (t, b) in enumerate(kinds):
    x = Inches(0.5 + i * 4.2)
    add_round(s, x, Inches(3.1), Inches(4.0), Inches(3.35), CARD, LINE)
    tb(s, x + Inches(0.22), Inches(3.3), Inches(3.55), Inches(0.45), t, 16, CRIMSON, True)
    tb(s, x + Inches(0.22), Inches(3.85), Inches(3.55), Inches(2.35), b, 14, INK2)
footer(s, 21)

# ============================================================
# SLIDE 22  空白意义
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "05  空白手法", "空白如何生成意义：从欲望中介到命运思考")
steps2 = [
    ("不写内心", "她不像真人，而像欲望之神、命运之神。非人/神性的一面被空白托出。"),
    ("三男之死", "她只是导火索。表面煞气，实则每人被自己的欲望吞没。吴家阿婆的妲己之喻，暴露的是徐壮图内心的外化。"),
    ("公馆乌托邦", "纸醉金迷麻痹「离家的孩子」。她以近乎上帝的视角旁观厮杀，繁华被推向虚无。"),
    ("麻将神谕", "「一个人的命运往往不受控制」——牌局说的是人生。欲望由人而生，人摆脱不了欲望，这本身即是命运。"),
]
for i, (t, b) in enumerate(steps2):
    y = Inches(1.48 + i * 1.25)
    add_round(s, Inches(0.5), y, Inches(12.3), Inches(1.12), CARD, LINE)
    tb(s, Inches(0.75), y + Inches(0.12), Inches(0.55), Inches(0.85), f"0{i+1}", 18, GOLD, True, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, Inches(1.5), y + Inches(0.15), Inches(2.3), Inches(0.8), t, 16, CRIMSON, True, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, Inches(3.9), y + Inches(0.18), Inches(8.6), Inches(0.8), b, 14, INK2, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 22)

# ============================================================
# SLIDE 23  意象总览
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "06  环境 · 意象 · 象征", "先分清三个概念，再看本篇如何叠用")
defs = [
    ("意象", "可感知的物象，承载情思。", "晚香玉、麻将、银白旗袍、郁金香、杏仁豆腐上的樱桃"),
    ("隐喻", "以此喻彼，使无生命者有感觉情欲（维柯）。", "麻将桌 = 人生；女祭司 = 送人上祭坛；晚香玉花语 = 危险的快乐"),
    ("象征", "具体形象指向超越自身的抽象意义，可多层。", "尹雪艳 → 欲望/命运/时间/旧上海；白色 → 死亡；红色 → 血光"),
    ("环境", "人物活动的空间，反作用于人。", "尹公馆：暖炉冷气、霞飞路排场——海市蜃楼，创造并囚禁这群人"),
]
for i, (t, d, e) in enumerate(defs):
    y = Inches(1.48 + i * 1.25)
    add_round(s, Inches(0.5), y, Inches(12.3), Inches(1.12), CARD, LINE)
    tb(s, Inches(0.75), y + Inches(0.32), Inches(1.6), Inches(0.5), t, 18, CRIMSON, True)
    tb(s, Inches(2.5), y + Inches(0.15), Inches(4.0), Inches(0.85), d, 13, MUTED, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, Inches(6.6), y + Inches(0.15), Inches(5.9), Inches(0.85), e, 13, INK2, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 23)

# ============================================================
# SLIDE 24  白与红
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "06  环境 · 意象 · 象征", "色彩隐喻：白色与红色互相对照")
add_round(s, Inches(0.5), Inches(1.5), Inches(6.1), Inches(5.0), CARD, LINE)
tb(s, Inches(0.75), Inches(1.7), Inches(5.6), Inches(0.4), "白  ·  宁静、圣洁、死亡", 17, RGBColor(0x5A, 0x6A, 0x78), True)
whites = [
    "名字之「雪」、雪白肌肤、素白旗袍",
    "玉梨花、冰雪精灵、银白女祭司",
    "月白旗袍与绣花鞋、银坠子",
    "雪白喷花露水的冰面巾",
    "殡仪馆「花圈丧幛白簇簇」",
    "葬礼上她仍素白打扮",
]
tb_multi(s, Inches(0.75), Inches(2.25), Inches(5.6), Inches(4.0),
         [{"text": "·  " + x, "size": 15} for x in whites])
add_round(s, Inches(6.85), Inches(1.5), Inches(6.0), Inches(5.0), CARD, LINE)
tb(s, Inches(7.1), Inches(1.7), Inches(5.5), Inches(0.4), "红  ·  欲望、血腥、预兆", 17, CRIMSON, True)
reds = [
    "名字之「艳」；血红郁金香破例上鬓",
    "桃花心红木桌椅（公馆的欲望底色）",
    "杏仁豆腐上「却」放两颗鲜红樱桃",
    "吴经理眼圈溃烂的粉红肉（衰老之红）",
    "「我来吃你的红」（彩头 / 血）",
    "重煞见血：王、洪、徐的血光之路",
]
tb_multi(s, Inches(7.1), Inches(2.25), Inches(5.5), Inches(4.0),
         [{"text": "·  " + x, "size": 15} for x in reds])
footer(s, 24)

# ============================================================
# SLIDE 25  关键意象
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "06  环境 · 意象 · 象征", "四个关键意象：晚香玉 · 麻将 · 公馆 · 女祭司")
imgs = [
    ("晚香玉", "室内本不宜放置的浓香，文中出现三次。徐壮图嗅到甜香，离开时花仍吐浓香。花语：危险的快乐。尹雪艳本身如晚香玉——由纯白转为重煞，再化为白色花圈。"),
    ("麻将", "社交娱乐，也是人生缩影。「命运往往不受控制」。追求胜欲、结果未知；牌可再打，人生迷失无回路。她的口才如神谕，补给斗志，却改变不了过眼云烟。"),
    ("尹公馆", "冬天暖炉、夏天冷气，维持霞飞路排场。作废头衔被娇声称呼，优越感如诰封。海市蜃楼、精神药剂、游子的乌托邦——也是屠宰场式的牌桌。"),
    ("女祭司", "巡视牌桌「通身银白」。暗示她将把这群人一个个送上祭坛。宗教意象把社交场面升格为祭祀与死亡仪式。"),
]
for i, (t, b) in enumerate(imgs):
    col, row = i % 2, i // 2
    x = Inches(0.45 + col * 6.4)
    y = Inches(1.48 + row * 2.55)
    add_round(s, x, y, Inches(6.2), Inches(2.4), CARD, LINE)
    tb(s, x + Inches(0.25), y + Inches(0.18), Inches(5.7), Inches(0.4), t, 17, CRIMSON, True)
    tb(s, x + Inches(0.25), y + Inches(0.65), Inches(5.7), Inches(1.55), b, 13, INK2)
footer(s, 25)

# ============================================================
# SLIDE 26  语言特色
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "07  语言特色", "客观冷静的外壳，里面是暗示、古典与张力")
langs = [
    ("客观与冷静", "像局外人记录。表面无感情，内里炽烈：对客人谴责仍带同情——可憎、可怜、可悲。"),
    ("古典词语", "体面、相干、台子、诰封……古朴清新带沧桑，与家族衰败、旧朝记忆同调。"),
    ("暗示性", "煞星应验；女祭司；「尹雪艳」一名：雪=内冷善变，艳=外美手段。神秘阴影笼罩全文。"),
    ("比喻", "玉梨花、冰雪精灵、三月微风、万年青。把不可见的魅力写成可感的形象。"),
    ("意象化描写", "晚香玉、服饰、麻将、挖花，把公馆写成诗意而抽象的画面，推向思辨。"),
    ("对话与词汇", "对话见性格关系；「熏住」「迎过来」「冒出火」把迷恋写到动作与感官。"),
]
for i, (t, b) in enumerate(langs):
    col, row = i % 3, i // 3
    x = Inches(0.45 + col * 4.25)
    y = Inches(1.5 + row * 2.5)
    add_round(s, x, y, Inches(4.05), Inches(2.3), CARD, LINE)
    tb(s, x + Inches(0.2), y + Inches(0.18), Inches(3.65), Inches(0.4), t, 15, CRIMSON, True)
    tb(s, x + Inches(0.2), y + Inches(0.65), Inches(3.65), Inches(1.45), b, 13, INK2)
footer(s, 26)

# ============================================================
# SLIDE 27  反讽结构
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "07  语言特色与反讽", "反讽是贯穿全文的结构原则，不只是几句挖苦")
tb(s, Inches(0.55), Inches(1.42), Inches(12.2), Inches(0.4),
   "概念：表面意思与真实意思分裂；读者须克服叙述者，才能到达隐含作者。", 14, MUTED)
irons = [
    ("话语反讽", "「老当益壮」感叹号很坚定，上文却是风湿、溃烂的眼。意义被语境扭曲；吴经理听了却「心中熨帖」。"),
    ("情景反讽", "大悲（葬礼）接大喜（四大喜）；她不悲不喜。「吃你的红」点破否极泰来的痴想，对方浑然不知。"),
    ("人物反讽", "以人的身体存世，却不老不死、高位俯瞰。以虚写实、以静衬动，反讽人世厮杀。"),
    ("结构反讽", "表面：永恒人物重现京沪繁华。实际：时过境迁，人们仍执迷追逐死亡之力，用享乐逃避无常。"),
]
for i, (t, b) in enumerate(irons):
    col, row = i % 2, i // 2
    x = Inches(0.5 + col * 6.4)
    y = Inches(1.95 + row * 2.25)
    add_round(s, x, y, Inches(6.15), Inches(2.1), CARD, LINE)
    tb(s, x + Inches(0.25), y + Inches(0.2), Inches(5.7), Inches(0.4), t, 16, CRIMSON, True)
    tb(s, x + Inches(0.25), y + Inches(0.7), Inches(5.7), Inches(1.2), b, 14, INK2)
footer(s, 27)

# ============================================================
# SLIDE 28  反讽三对照
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "07  语言特色与反讽", "三组对照把反讽写进结构")
triples2 = [
    ("容颜：流逝 / 不逝", "宋太太痴肥气喘，吴经理白发溃眼；她「像枝万年青，愈来愈年青」。\n用没有止境的永生，反讽人的衰老与死亡。\n「人生代代无穷已，江月年年望相似。」"),
    ("态度：热烈 / 清冷", "王、洪、徐及满座新旧客人着魔追逐；她停一宵、带走家当、鞠躬握手后当晚开局。\n他们需要这味精神药剂，重煞也抵不过逃避的兴味。"),
    ("欲望 / 死亡", "公馆给人身份认同，填不满的欲望让人甘之如饴奔向冷漠。\n牌局如人生。她加速纵欲者的死亡。永生与衰老、不悲不喜与为情所困——命运短暂而可笑。"),
]
for i, (t, b) in enumerate(triples2):
    x = Inches(0.45 + i * 4.25)
    add_round(s, x, Inches(1.5), Inches(4.1), Inches(4.95), CARD, LINE)
    add_rect(s, x, Inches(1.5), Inches(4.1), Inches(0.08), CRIMSON if i != 1 else GOLD)
    tb(s, x + Inches(0.22), Inches(1.75), Inches(3.7), Inches(0.7), t, 16, CRIMSON, True)
    tb(s, x + Inches(0.22), Inches(2.55), Inches(3.7), Inches(3.6), b, 14, INK2)
footer(s, 28)

# ============================================================
# SLIDE 29  主旨六层
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "08  主旨", "六层主旨：由时间命运到欲望虚无（答题选切题的两三重写）")
themes = [
    ("时间与命运\n不可抗拒", "不老 vs 衰老死亡。王、洪、徐都无法逃脱摆布。"),
    ("繁华与衰落\n的对比", "百乐门 / 尹公馆 vs 台北的今不如昔。公馆是荣光的残影。"),
    ("人性的复杂\n与矛盾", "煞气多是外界投射；她从容生活，却被读成祸水。向往与恐惧并存。"),
    ("怀旧与现实\n的冲突", "身在当下、心在往昔。公馆是避风港，也是不肯醒来的证明。"),
    ("女性的力量\n与悲剧", "迷倒男女，自己仍依附男性生存；被污名，也被需要。"),
    ("轮回、欲望\n与无常", "追逐故事一直没变。欲望之神不老，人是命运手中的玩偶。"),
]
for i, (t, b) in enumerate(themes):
    col, row = i % 3, i // 3
    x = Inches(0.45 + col * 4.25)
    y = Inches(1.48 + row * 2.55)
    add_round(s, x, y, Inches(4.1), Inches(2.4), CARD, LINE)
    tb(s, x + Inches(0.22), y + Inches(0.15), Inches(0.5), Inches(0.4), f"{i+1}", 16, GOLD, True)
    tb(s, x + Inches(0.7), y + Inches(0.15), Inches(3.15), Inches(0.85), t, 15, CRIMSON, True)
    tb(s, x + Inches(0.22), y + Inches(1.1), Inches(3.65), Inches(1.1), b, 13, INK2)
footer(s, 29)

# ============================================================
# SLIDE 30  手法通向主旨
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "08  主旨", "手法如何服务主旨：一条可写进论文的因果链")
chains = [
    ("塑造：银白 + 不写内心 + 细节之冷", "→ 人物非人化", "才能同时承担欲望、命运、死神"),
    ("叙事：全知不入内 + 倾慕口吻", "→ 不可靠叙述", "读者经历「着迷→察觉裂隙→批判」"),
    ("插叙：上海插入台北", "→ 今昔对照", "怀旧是自欺，繁华是仿制品"),
    ("空白：失声、态度词、环境满而人空", "→ 象征敞开", "意义在读者填补中多重实现"),
    ("色彩/晚香玉/麻将", "→ 感官与寓言", "危险的快乐、命运不受控"),
    ("反讽贯穿语调与结构", "→ 主题深化", "愚顽自欺无法摆脱生死"),
]
add_round(s, Inches(0.45), Inches(1.42), Inches(12.4), Inches(5.25), CARD, LINE)
tb(s, Inches(0.7), Inches(1.55), Inches(5.0), Inches(0.35), "手法组合", 12, CRIMSON, True)
tb(s, Inches(6.0), Inches(1.55), Inches(2.8), Inches(0.35), "直接效果", 12, CRIMSON, True)
tb(s, Inches(8.9), Inches(1.55), Inches(3.6), Inches(0.35), "通向的主旨", 12, CRIMSON, True)
add_rect(s, Inches(0.7), Inches(1.95), Inches(11.9), Inches(0.015), LINE)
for i, (a, b, c) in enumerate(chains):
    y = Inches(2.1 + i * 0.7)
    tb(s, Inches(0.7), y, Inches(5.1), Inches(0.6), a, 13, INK, True)
    tb(s, Inches(6.0), y, Inches(2.8), Inches(0.6), b, 13, GOLD, True)
    tb(s, Inches(8.9), y, Inches(3.7), Inches(0.6), c, 13, INK2)
footer(s, 30)

# ============================================================
# SLIDE 31  答题路径
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "08  主旨与答题路径", "IB / 文学论述：把概念用对，把例子写透")
tips = [
    ("审题先定层", "问「人物」就走形象+塑造；问「叙事」走视角与不可靠；问「主题」必须有手法作证据，不要只复述情节。"),
    ("例子要带原文词", "「总也不老」「停了一宵」「还算有良心」「吃你的红」\n「悲天悯人」「却放着两颗鲜红的樱桃」——词就是论点的钉子。"),
    ("反讽必写裂隙", "说出「谁在说」与「谁在想」：叙述者 / 隐含作者 / 读者。没有这三方，反讽只是形容词。"),
    ("象征不要清单化", "一次作文选 1–2 个象征轴写深（欲望或时间），用色彩、空白、情节循环去支撑，比罗列六种象征更有力。"),
    ("评价可落在张力", "神/妖、冷/热、白/红、永远/无常。点明张力如何让「台北人」的怀旧显出可笑与可哀。"),
    ("常见误读", "停在赞美佳人；或简单骂她祸水。两种都是被叙述者或吴家阿婆牵着走。文本要的是欲望与命运的思考。"),
]
for i, (t, b) in enumerate(tips):
    col, row = i % 2, i // 2
    x = Inches(0.45 + col * 6.4)
    y = Inches(1.45 + row * 1.7)
    add_round(s, x, y, Inches(6.2), Inches(1.55), CARD, LINE)
    tb(s, x + Inches(0.22), y + Inches(0.12), Inches(5.8), Inches(0.35), t, 15, CRIMSON, True)
    tb(s, x + Inches(0.22), y + Inches(0.5), Inches(5.8), Inches(0.9), b, 12, INK2)
footer(s, 31)

# ============================================================
# SLIDE 32  复习清单
# ============================================================
s = prs.slides.add_slide(blank)
bg(s)
section_bar(s, "复习清单", "合上笔记前，能用自己的话回答这些")
qs = [
    "「尹雪艳总也不老」为何既是人物设定，又是反讽基调？",
    "「雪」与「艳」如何在服饰、色彩、行为上同时展开？",
    "以形写神、白描、细节三者，各举一例并说明效果有何不同。",
    "全知但不入内心，怎样制造神秘，又怎样为嘲讽服务？",
    "叙述者说「还算有良心」，隐含作者可能想让读者感到什么？",
    "空白（失声、态度词、环境）如何把她从「女人」推向「欲望/命运」？",
    "晚香玉三次出现，如何把「危险的快乐」写进徐壮图的路线？",
    "白与红在葬礼/寿宴/樱桃/吃红四处，怎样把死亡主题视觉化？",
    "插叙让上海进入台北，对「怀旧」主题做了什么？",
    "若只选一个贯穿手法统领全篇，为什么是反讽而不是象征？",
]
for i, q in enumerate(qs):
    col, row = i % 2, i // 2
    x = Inches(0.45 + col * 6.4)
    y = Inches(1.45 + row * 1.0)
    n = f"{i+1:02d}"
    add_round(s, x, y, Inches(6.2), Inches(0.88), CARD, LINE)
    tb(s, x + Inches(0.15), y + Inches(0.22), Inches(0.55), Inches(0.45), n, 14, GOLD, True)
    tb(s, x + Inches(0.7), y + Inches(0.12), Inches(5.3), Inches(0.65), q, 13, INK2, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 32)

# ============================================================
# SLIDE 33  结束
# ============================================================
s = prs.slides.add_slide(blank)
bg(s, WINE)
add_rect(s, 0, 0, Inches(0.18), H, GOLD)
tb(s, Inches(0.8), Inches(2.0), Inches(12), Inches(0.4), "收束", 14, GOLD, True)
tb(s, Inches(0.8), Inches(2.5), Inches(11.5), Inches(1.2),
   "永远的，不是美人，\n而是欲望本身，以及人对幻觉的需要。", 28, WHITE, True)
tb(s, Inches(0.8), Inches(4.3), Inches(11.5), Inches(1.3),
   "尹雪艳总也不老——因为心魔不老。\n公馆总还开着牌局——因为自欺总还需要一个舞台。\n文学手法的意义，是让读者在赞美的语调里，听见死亡的寒意。",
   16, RGBColor(0xE8, 0xC8, 0xC4))
tb(s, Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.4),
   "白先勇  ·  《永远的尹雪艳》  ·  《台北人》", 14, GOLD)

out_cn = "/workspace/永远的尹雪艳_文学概念与手法.pptx"
out_en = "/workspace/Forever-Yin-Xueyan.pptx"
prs.save(out_cn)
prs.save(out_en)
print("saved", out_cn, "slides", len(prs.slides))
print("saved", out_en)
