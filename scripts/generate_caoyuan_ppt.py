#!/usr/bin/env python3
"""Generate MYP lesson PPT for 统编版六年级上册《草原》."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

INK = RGBColor(0x1A, 0x2F, 0x28)
TEAL = RGBColor(0x1F, 0x5C, 0x4A)
GRASS = RGBColor(0x3D, 0x8B, 0x5A)
SKY = RGBColor(0x2F, 0x6F, 0x8F)
SUN = RGBColor(0xC4, 0x8A, 0x1A)
MIST = RGBColor(0xE6, 0xF2, 0xEA)
PAPER = RGBColor(0xF5, 0xFA, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0x3A, 0x4F, 0x45)

TOTAL = 12


def set_run(run, size=18, bold=False, color=INK, name="微软雅黑"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", name)


def add_bg(slide, color):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_bar(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text(slide, left, top, width, height, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_paras(slide, left, top, width, height, lines, size=16, color=INK, bold_first=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = line
        set_run(run, size=size, bold=(bold_first and i == 0), color=color)
    return box


def new_slide(bg=PAPER):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, bg)
    return slide


def footer(slide, page):
    add_text(slide, Inches(0.6), Inches(7.05), Inches(10), Inches(0.35),
             "MYP 语言与文学 · 《草原》· 全球背景：身份认同与人际关系", size=11, color=SOFT)
    add_text(slide, Inches(11.5), Inches(7.05), Inches(1.3), Inches(0.35),
             f"{page}/{TOTAL}", size=11, color=SOFT, align=PP_ALIGN.RIGHT)


# 1 Cover
s = new_slide(MIST)
add_bar(s, Inches(0), Inches(0), Inches(0.35), prs.slide_height, TEAL)
add_text(s, Inches(1.0), Inches(1.5), Inches(11), Inches(1.0), "草原", size=60, bold=True, color=TEAL)
add_text(s, Inches(1.0), Inches(2.7), Inches(11), Inches(0.6),
         "地方与相遇，如何让我们看见彼此？", size=26, color=INK)
add_text(s, Inches(1.0), Inches(3.6), Inches(11), Inches(0.8),
         "统编版六年级上册 · 老舍\nMYP Year 1 · 语言与文学 · 2 课时教学课件", size=16, color=SOFT)
add_text(s, Inches(1.0), Inches(5.2), Inches(11), Inches(0.6),
         "全球背景：身份认同与人际关系　｜　关键概念：联系", size=16, bold=True, color=GRASS)
footer(s, 1)

# 2 Framework
s = new_slide()
add_bar(s, Inches(0), Inches(0), prs.slide_width, Inches(0.15), TEAL)
add_text(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.5), "MYP 探究框架", size=28, bold=True, color=TEAL)
rows = [
    ("全球背景", "身份认同与人际关系 —— 归属、友谊、跨文化理解"),
    ("关键概念", "联系（Connections）"),
    ("相关概念", "场景 · 风格 · 主题"),
    ("探究声明", "文学作品通过描绘地方与相遇，揭示文化身份如何在联系中被理解与深化。"),
    ("ATL 技能", "沟通 · 思考 · 社交（换位思考）· 研究（提取信息）"),
]
y = 1.1
for label, val in rows:
    add_bar(s, Inches(0.6), Inches(y), Inches(2.3), Inches(0.85), MIST)
    add_text(s, Inches(0.75), Inches(y + 0.25), Inches(2.0), Inches(0.4), label, size=15, bold=True, color=TEAL)
    add_text(s, Inches(3.1), Inches(y + 0.2), Inches(9.5), Inches(0.55), val, size=15, color=INK)
    y += 1.0
footer(s, 2)

# 3 Inquiry questions
s = new_slide()
add_bar(s, Inches(0), Inches(0), prs.slide_width, Inches(0.15), TEAL)
add_text(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.5), "三条探究问题", size=28, bold=True, color=TEAL)
qs = [
    ("事实性", "课文写了哪些风景？按什么顺序写了哪些人事活动？", GRASS),
    ("概念性", "作者如何把「景美」与「情美」联系起来？地方如何影响人的感受与关系？", SKY),
    ("辩论性", "真正的友谊是否一定需要共同的语言与相同的文化背景？", SUN),
]
y = 1.2
for title, q, c in qs:
    add_bar(s, Inches(0.7), Inches(y), Inches(12), Inches(1.5), MIST)
    add_bar(s, Inches(0.7), Inches(y), Inches(0.2), Inches(1.5), c)
    add_text(s, Inches(1.2), Inches(y + 0.25), Inches(11), Inches(0.4), title, size=16, bold=True, color=c)
    add_text(s, Inches(1.2), Inches(y + 0.7), Inches(11), Inches(0.5), q, size=18, color=INK)
    y += 1.7
footer(s, 3)

# 4 Goals
s = new_slide()
add_bar(s, Inches(0), Inches(0), prs.slide_width, Inches(0.15), TEAL)
add_text(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.5), "学习目标与成功标准", size=28, bold=True, color=TEAL)
add_paras(s, Inches(0.7), Inches(1.1), Inches(5.8), Inches(5), [
    "【学习目标】",
    "1. 理清：风光 → 迎客—相见—款待—联欢—话别",
    "2. 理解概念：联系、场景、风格、主题",
    "3. 品味比喻、拟人与情景交融",
    "4. 写一次「温暖相遇」短文（100–150字）",
    "5. 用开放心态理解跨文化交往",
], size=15, bold_first=True)
add_paras(s, Inches(6.9), Inches(1.1), Inches(5.8), Inches(5), [
    "【我能做到】",
    "· 概括课文的「景」与「情」",
    "· 找出 ≥3 处美/热/亲的语句并说明方法",
    "· 解释结尾诗句与全文的联系",
    "· 短文有动作细节与心情变化",
    "· 用一句话回应：联系如何改变我们对他人的看法",
], size=15, bold_first=True)
footer(s, 4)

# 5 Lesson 1 overview
s = new_slide()
add_bar(s, Inches(0), Inches(0), prs.slide_width, Inches(0.15), TEAL)
add_text(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.5), "课时一：地方如何塑造感受", size=28, bold=True, color=TEAL)
add_text(s, Inches(0.7), Inches(1.0), Inches(12), Inches(0.4),
         "读出「景美」，看见「联系」的起点", size=18, color=GRASS)
steps = [
    ("0–8′", "Tuning in", "两张图一个问题：地方如何影响心情与待人方式？"),
    ("8–23′", "Finding out", "初读 + 双栏笔记（景 / 情）+ 结构拼图"),
    ("23–38′", "Sorting out", "风光段三色标注：颜色形态 / 修辞 / 心情"),
    ("38–45′", "Reflection", "若草原会说话……；仿写一句家乡/校园景色"),
]
y = 1.6
for t, phase, desc in steps:
    add_bar(s, Inches(0.6), Inches(y), Inches(1.5), Inches(0.95), GRASS)
    add_text(s, Inches(0.7), Inches(y + 0.3), Inches(1.3), Inches(0.4), t, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, Inches(2.3), Inches(y + 0.15), Inches(3.2), Inches(0.35), phase, size=16, bold=True, color=TEAL)
    add_text(s, Inches(2.3), Inches(y + 0.5), Inches(10), Inches(0.4), desc, size=15, color=INK)
    y += 1.15
footer(s, 5)

# 6 Dual column
s = new_slide()
add_bar(s, Inches(0), Inches(0), prs.slide_width, Inches(0.15), TEAL)
add_text(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.5), "初读工具：双栏笔记", size=28, bold=True, color=TEAL)
add_bar(s, Inches(0.7), Inches(1.2), Inches(5.7), Inches(5.0), MIST)
add_bar(s, Inches(6.9), Inches(1.2), Inches(5.7), Inches(5.0), MIST)
add_text(s, Inches(0.95), Inches(1.4), Inches(5.2), Inches(0.5), "我看见的「地方」（景）", size=18, bold=True, color=TEAL)
add_text(s, Inches(7.15), Inches(1.4), Inches(5.2), Inches(0.5), "我看见的「人与关系」（情）", size=18, bold=True, color=TEAL)
add_paras(s, Inches(0.95), Inches(2.2), Inches(5.2), Inches(3.5), [
    "· 天、空气、小丘、绿毯",
    "· 羊群、骏马、大牛",
    "· 一碧千里……",
    "",
    "提示：先记下画面，再想心情",
], size=16)
add_paras(s, Inches(7.15), Inches(2.2), Inches(5.2), Inches(3.5), [
    "· 谁来迎？怎么迎？",
    "· 怎样相见、款待、联欢？",
    "· 话别时人们怎样不舍？",
    "",
    "提示：关注关系如何变化",
], size=16)
footer(s, 6)

# 7 Style focus
s = new_slide()
add_bar(s, Inches(0), Inches(0), prs.slide_width, Inches(0.15), TEAL)
add_text(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.5), "精读：风格如何写出「地方感」", size=28, bold=True, color=TEAL)
add_paras(s, Inches(0.7), Inches(1.1), Inches(12), Inches(5.2), [
    "概念性问题：作者用哪些语言策略，让读者「进入」草原？",
    "",
    "关注（请对照课文原文）：",
    "· 「那里的天比别处的天更可爱……」——比较中见独特",
    "· 「一碧千里，而并不茫茫」——辽阔却有边界感/亲切感",
    "· 羊群像绣在绿毯上的白花 —— 比喻让静止的画面「活」",
    "· 「翠色欲流……流入云际」——颜色像有了动作",
    "· 骏马大牛静立「回味」——拟人写出宁静与诗意",
    "",
    "活动：三色标注 = 颜色形态 / 修辞 / 心情　→　追问「景物与心情如何联系？」",
], size=16, bold_first=True)
footer(s, 7)

# 8 Lesson 2
s = new_slide()
add_bar(s, Inches(0), Inches(0), prs.slide_width, Inches(0.15), TEAL)
add_text(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.5), "课时二：相遇如何建立联系", size=28, bold=True, color=TEAL)
add_text(s, Inches(0.7), Inches(1.0), Inches(12), Inches(0.4),
         "读出「情美」，迁移到真实世界", size=18, color=GRASS)
steps2 = [
    ("0–5′", "回接", "从「风景中的我」走向「关系中的我们」"),
    ("5–17′", "场面阅读", "迎—见—待—欢—别：动作、心情、联系"),
    ("17–27′", "主旨深挖", "结尾诗句三层追问 + 迷你辩论"),
    ("27–40′", "写作迁移", "《一次温暖的相遇》100–150字"),
    ("40–45′", "反思行动", "三问反思 + 采访一个跨群体故事"),
]
y = 1.55
for t, phase, desc in steps2:
    add_bar(s, Inches(0.6), Inches(y), Inches(1.5), Inches(0.85), SKY)
    add_text(s, Inches(0.7), Inches(y + 0.25), Inches(1.3), Inches(0.4), t, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, Inches(2.3), Inches(y + 0.1), Inches(2.5), Inches(0.35), phase, size=16, bold=True, color=TEAL)
    add_text(s, Inches(5.0), Inches(y + 0.2), Inches(7.5), Inches(0.45), desc, size=15, color=INK)
    y += 0.95
footer(s, 8)

# 9 Global context deep dive
s = new_slide()
add_bar(s, Inches(0), Inches(0), prs.slide_width, Inches(0.15), TEAL)
add_text(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.5), "全球背景渗透：三层追问", size=28, bold=True, color=TEAL)
add_text(s, Inches(0.7), Inches(1.0), Inches(12), Inches(0.45),
         "聚焦：「蒙汉情深何忍别，天涯碧草话斜阳」", size=18, bold=True, color=GRASS)
layers = [
    ("字面层", "谁不忍别？在什么地方话别？"),
    ("文本层", "为什么「碧草斜阳」必须与「蒙汉情深」写在一起？"),
    ("全球背景层", "跨越语言、习俗、地域的友谊靠什么建立？可迁移到哪些当代情境？"),
]
y = 1.7
for title, q in layers:
    add_bar(s, Inches(0.7), Inches(y), Inches(12), Inches(1.2), MIST)
    add_text(s, Inches(1.0), Inches(y + 0.2), Inches(11.4), Inches(0.35), title, size=16, bold=True, color=TEAL)
    add_text(s, Inches(1.0), Inches(y + 0.6), Inches(11.4), Inches(0.4), q, size=17, color=INK)
    y += 1.4
footer(s, 9)

# 10 Debate + writing
s = new_slide()
add_bar(s, Inches(0), Inches(0), prs.slide_width, Inches(0.15), TEAL)
add_text(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.5), "辩论 + 写作迁移", size=28, bold=True, color=TEAL)
add_bar(s, Inches(0.6), Inches(1.1), Inches(5.9), Inches(5.2), MIST)
add_bar(s, Inches(6.8), Inches(1.1), Inches(5.9), Inches(5.2), MIST)
add_text(s, Inches(0.85), Inches(1.3), Inches(5.4), Inches(0.4), "迷你辩论", size=18, bold=True, color=TEAL)
add_paras(s, Inches(0.85), Inches(1.9), Inches(5.4), Inches(4), [
    "正方：共同文化更易建立友谊",
    "反方：真诚行动比相同背景更重要",
    "",
    "规则：必须引用课文一个细节作证据",
    "",
    "目的：把「民族团结」落成",
    "可观察的行为与选择",
], size=15)
add_text(s, Inches(7.05), Inches(1.3), Inches(5.4), Inches(0.4), "短写句架", size=18, bold=True, color=TEAL)
add_paras(s, Inches(7.05), Inches(1.9), Inches(5.4), Inches(4), [
    "《一次温暖的相遇》100–150字",
    "",
    "1. 在____，我遇到了____。",
    "2. 起初我觉得____；",
    "   后来因为____（动作），我感到____。",
    "3. 这次相遇让我明白：",
    "   联系往往开始于____。",
], size=15)
footer(s, 10)

# 11 Assessment
s = new_slide()
add_bar(s, Inches(0), Inches(0), prs.slide_width, Inches(0.15), TEAL)
add_text(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.5), "评估与量规（微型）", size=28, bold=True, color=TEAL)
headers = ["维度", "达标", "良好", "优秀"]
rows_r = [
    ["组织", "有起因与结果", "相遇过程清楚", "关系变化清晰可见"],
    ["细节", "1 处动作", "2–3 处具体细节", "动作+心情+场合融合"],
    ["概念迁移", "提到友好", "点出联系/理解", "连接课文与现实"],
    ["语言", "基本通顺", "有生动动词", "合理迁移课文写法"],
]
xs = [0.7, 2.5, 5.5, 9.2]
widths = [1.6, 2.8, 3.5, 3.5]
for i, h in enumerate(headers):
    add_bar(s, Inches(xs[i]), Inches(1.2), Inches(widths[i]), Inches(0.5), TEAL)
    add_text(s, Inches(xs[i] + 0.1), Inches(1.28), Inches(widths[i] - 0.15), Inches(0.35),
             h, size=14, bold=True, color=WHITE)
y = 1.8
for r in rows_r:
    for i, cell in enumerate(r):
        add_bar(s, Inches(xs[i]), Inches(y), Inches(widths[i]), Inches(0.9), MIST if i == 0 else PAPER)
        add_text(s, Inches(xs[i] + 0.1), Inches(y + 0.28), Inches(widths[i] - 0.15), Inches(0.45),
                 cell, size=13, color=INK)
    y += 0.95
add_text(s, Inches(0.7), Inches(5.7), Inches(12), Inches(0.8),
         "可对接 MYP：A 分析 · B 组织 · C 创作文本 · D 使用语言（本课以形成性 + 微型总结性为主）",
         size=14, color=SOFT)
footer(s, 11)

# 12 Close
s = new_slide(MIST)
add_bar(s, Inches(0), Inches(0), Inches(0.35), prs.slide_height, TEAL)
add_text(s, Inches(1.0), Inches(1.6), Inches(11), Inches(0.7), "带着概念离开教室", size=32, bold=True, color=TEAL)
add_paras(s, Inches(1.0), Inches(2.6), Inches(11), Inches(3), [
    "今天请记住三句话：",
    "1. 草原不仅是风景，也是关系发生的地方。",
    "2. 「不同」可以通过具体行动变成「亲近」。",
    "3. 读写的目标，是更清楚地理解自己与他人。",
    "",
    "课后行动：采访一个「跨群体互助」的真实故事，用三句话分享。",
], size=18)
footer(s, 12)

out = "/workspace/ppt/草原_MYP教案课件.pptx"
prs.save(out)
print("saved", out, "slides", len(prs.slides))
