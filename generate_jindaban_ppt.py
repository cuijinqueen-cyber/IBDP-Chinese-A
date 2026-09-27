#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""《金大班的最后一夜》精读教学课件（据课堂笔记＋文本引证）"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

BG = RGBColor(0x1C, 0x14, 0x18)
CARD = RGBColor(0x2A, 0x1F, 0x26)
ACCENT = RGBColor(0xC9, 0x8B, 0x3C)      # 金
ACCENT2 = RGBColor(0xB8, 0x4A, 0x5A)     # 夜巴黎红
TEXT = RGBColor(0xF5, 0xEE, 0xE6)
MUTED = RGBColor(0xB0, 0xA4, 0xA8)
QUOTE = RGBColor(0xE8, 0xC9, 0x8A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W, H = prs.slide_width, prs.slide_height


def set_run(run, size=18, bold=False, color=TEXT, font="Microsoft YaHei"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", font)


def add_bg(slide, color=BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_bar(slide, left, top, width, height, color=ACCENT):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_card(slide, left, top, width, height, color=CARD):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_textbox(slide, left, top, width, height, text, size=18, bold=False,
                color=TEXT, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_paras(slide, left, top, width, height, lines, size=14, color=TEXT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            text, opts = item
        else:
            text, opts = item, {}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(opts.get("space_after", 6))
        run = p.add_run()
        run.text = text
        set_run(
            run,
            size=opts.get("size", size),
            bold=opts.get("bold", False),
            color=opts.get("color", color),
        )
    return box


def section_header(slide, section_no, title, subtitle=""):
    add_bg(slide)
    add_bar(slide, Inches(0), Inches(0), Inches(0.18), H, ACCENT2)
    add_textbox(slide, Inches(0.55), Inches(0.3), Inches(12), Inches(0.35),
                f"《金大班的最后一夜》精读  ·  {section_no}", size=13, color=ACCENT)
    add_textbox(slide, Inches(0.55), Inches(0.65), Inches(12), Inches(0.55),
                title, size=28, bold=True, color=WHITE)
    if subtitle:
        add_textbox(slide, Inches(0.55), Inches(1.25), Inches(12), Inches(0.4),
                    subtitle, size=14, color=MUTED)


def quote_block(slide, left, top, width, height, quote, note=""):
    add_card(slide, left, top, width, height)
    add_bar(slide, left, top, Inches(0.08), height, ACCENT)
    lines = [(f"「{quote}」", {"size": 13, "color": QUOTE, "bold": True, "space_after": 6})]
    if note:
        lines.append((note, {"size": 12, "color": MUTED, "space_after": 0}))
    add_paras(slide, left + Inches(0.22), top + Inches(0.15),
              width - Inches(0.35), height - Inches(0.25), lines)


# ========== 封面 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_bar(slide, 0, 0, W, Inches(0.12), ACCENT)
add_textbox(slide, Inches(0.8), Inches(1.7), Inches(11), Inches(0.4),
            "白先勇 · 《台北人》", size=18, color=ACCENT)
add_textbox(slide, Inches(0.8), Inches(2.25), Inches(12), Inches(0.9),
            "《金大班的最后一夜》", size=44, bold=True, color=WHITE)
add_textbox(slide, Inches(0.8), Inches(3.3), Inches(11), Inches(0.45),
            "精读教学课件  ·  结合笔记 · 文本细节引证", size=20, color=MUTED)
add_paras(slide, Inches(0.8), Inches(4.3), Inches(11.5), Inches(2), [
    ("分析顺序", {"size": 14, "bold": True, "color": ACCENT, "space_after": 10}),
    ("一、人物形象  →  二、艺术手法  →  三、小说结构  →  四、环境",
     {"size": 16, "color": TEXT, "space_after": 12}),
    ("核心关键词：矛盾立体 · 今昔对照 · 一夜框一生 · 百乐门／夜巴黎",
     {"size": 14, "color": MUTED, "space_after": 0}),
])

# ========== 目录 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "导读", "目录", "据笔记四大板块展开，每页尽量附原文引证")
items = [
    ("01", "人物形象", "复杂立体多面：道路选择／处世态度／爱情态度"),
    ("02", "艺术手法", "对比衬托 · 反讽 · 白描 · 心理 · 语言动作 · 象征"),
    ("03", "小说结构", "一夜框一生 · 倒叙插叙 · 意识流 · 视角 · 开放结局"),
    ("04", "环境", "微观舞厅 · 历史背景 · 象征隐喻（舞厅／服饰／镜子）"),
]
for i, (no, tit, sub) in enumerate(items):
    y = Inches(2.0) + Inches(i * 1.15)
    add_card(slide, Inches(0.55), y, Inches(12.2), Inches(1.0))
    add_textbox(slide, Inches(0.85), y + Inches(0.2), Inches(1.2), Inches(0.5),
                no, size=24, bold=True, color=ACCENT)
    add_textbox(slide, Inches(2.3), y + Inches(0.15), Inches(10), Inches(0.35),
                tit, size=20, bold=True, color=WHITE)
    add_textbox(slide, Inches(2.3), y + Inches(0.52), Inches(10), Inches(0.35),
                sub, size=13, color=MUTED)

# ========== 一、人物总览 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "一 / 人物形象", "金兆丽：复杂 · 立体 · 多面",
               "笔记总括：内在矛盾性导致道路犹豫、处世善恶并存、爱情前后相反")
cards = [
    ("道路矛盾", "曾嘲姐妹「饿嫁」\n最终下嫁陈发荣\n孤傲 → 屈降依附"),
    ("处世矛盾", "救朱凤、赠钻戒\n妒任黛黛／怼童得怀\n既善又恶"),
    ("爱情矛盾", "为月如愿生孩子\n弃秦雄真心\n真情 → 物质依附"),
    ("悲剧位置", "舞场腐蚀中\n仍想保留求幸福自由\n二者难以调和"),
]
for i, (t, b) in enumerate(cards):
    x = Inches(0.45) + Inches(i * 3.2)
    add_card(slide, x, Inches(2.0), Inches(3.05), Inches(4.7))
    add_bar(slide, x, Inches(2.0), Inches(3.05), Inches(0.08), ACCENT if i == 0 else ACCENT2)
    add_textbox(slide, x + Inches(0.2), Inches(2.25), Inches(2.65), Inches(0.4),
                t, size=16, bold=True, color=ACCENT)
    add_paras(slide, x + Inches(0.2), Inches(2.85), Inches(2.65), Inches(3.5), [
        (line, {"size": 13, "color": TEXT if j == 0 else MUTED, "space_after": 6})
        for j, line in enumerate(b.split("\n"))
    ])

# ========== 道路矛盾 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "一 / 人物形象", "① 人生道路：独善其身 ↔ 屈降回归",
               "嘲「饿嫁」与最终下嫁——同一人的前后对翻")
quote_block(slide, Inches(0.5), Inches(1.95), Inches(6.05), Inches(2.35),
            "我才没有你们那样饿嫁，个个去捧块棺材板。",
            "引证 · 孤傲宣言：鄙薄姐妹回归家庭；争强好胜、我行我素")
quote_block(slide, Inches(6.8), Inches(1.95), Inches(6.05), Inches(2.35),
            "要一个像任黛黛那样的绸缎庄，当然要比她的那个大一倍，并且要开在富春楼的对面。",
            "引证 · 屈降自慰：物质目标合理化下嫁；妒羡与自我安慰并存")
add_card(slide, Inches(0.5), Inches(4.55), Inches(12.35), Inches(2.35))
add_paras(slide, Inches(0.8), Inches(4.75), Inches(11.8), Inches(2.0), [
    ("精读要点（据笔记）", {"size": 15, "bold": True, "color": ACCENT, "space_after": 8}),
    ("• 舞女非理想职业，嫁富商亦非梦想；贫寒逼入欢场二十年，年老色衰又逼入依附婚姻。",
     {"size": 13, "color": TEXT, "space_after": 5}),
    ("• 见陈发荣「还算真诚」便决定下嫁：不是爱上他，而是看清「黑暗无终点」只能依附才能活。",
     {"size": 13, "color": TEXT, "space_after": 5}),
    ("• 效果：贪图衣食无忧又厌倦失去自我——徘徊把内在矛盾推到台前。",
     {"size": 13, "color": MUTED, "space_after": 0}),
])

# ========== 处世矛盾 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "一 / 人物形象", "② 处世态度：同情弱小 ↔ 艳羡妒忌",
               "人格既高尚又微露卑下——舞场未磨尽善良，却日日腐蚀纯真")
quote_block(slide, Inches(0.5), Inches(1.95), Inches(6.05), Inches(2.5),
            "朱凤怀有身孕向她哭诉……臭骂一顿之后，又甩给她价值五百美金的一克拉半的大钻石戒指。",
            "引证 · 对弱小：挺身救朱凤、教「十八般武艺」、警戒勿被占身；不能自救却救助别人")
quote_block(slide, Inches(6.8), Inches(1.95), Inches(6.05), Inches(2.5),
            "筱红美「是一个头等难缠的刁妇，心黑手辣，耍了这些年就没见过她栽过跟头」。",
            "引证 · 对强者：妒羡赞叹＋尖酸怼童得怀；看透世事又被现实逼出小心眼")
add_card(slide, Inches(0.5), Inches(4.7), Inches(12.35), Inches(2.2))
add_paras(slide, Inches(0.8), Inches(4.9), Inches(11.8), Inches(1.8), [
    ("效果", {"size": 15, "bold": True, "color": ACCENT, "space_after": 8}),
    ("善与恶不是分期出现，而是同时交错：同情证明本性未死；妒忌证明环境已把人性拧弯。"
     "人物因此立体，而非「好舞女／坏舞女」扁平标签。",
     {"size": 14, "color": TEXT, "space_after": 0}),
])

# ========== 爱情矛盾 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "一 / 人物形象", "③ 爱情态度：真挚追求 ↔ 现实背弃",
               "月如＝心中圣洁区；秦雄真心等不起；陈发荣＝物质落脚点")
quote_block(slide, Inches(0.5), Inches(1.95), Inches(12.35), Inches(1.55),
            "那晚月如第一次到百乐门去，羞得连头都不抬起来……纯洁的像白纸一样。",
            "引证 · 真情：愿为他生孩子＝职业上自掘坟墓；堕胎后甚至想死——情根极深")
quote_block(slide, Inches(0.5), Inches(3.7), Inches(6.05), Inches(3.1),
            "秦雄拿出可怜巴巴的七万元存折……再积攒五年就能买房子讨她；她却无动于衷，下嫁前夜也没写信。",
            "背弃真爱：四十岁「没有功夫谈恋爱」，只要衣食无忧")
quote_block(slide, Inches(6.8), Inches(3.7), Inches(6.05), Inches(3.1),
            "见脸红害羞的年轻人，仍想到月如；与酷似月如者共舞时情绪低回。",
            "效果：表面背弃，内心未磨灭；灵肉撕裂在结尾回光一现")

# ========== 人物总括 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "一 / 人物形象", "人物总括：双重矛盾的结合体",
               "非简单线性变坏，而是道路／处世／爱情三线交叉的复杂真实")
add_card(slide, Inches(0.5), Inches(1.95), Inches(12.35), Inches(4.9))
add_paras(slide, Inches(0.85), Inches(2.2), Inches(11.7), Inches(4.4), [
    ("笔记核心判断", {"size": 16, "bold": True, "color": ACCENT, "space_after": 12}),
    ("金兆丽「集双重矛盾于一身」：现实残酷逼她审时度势，物质依赖维持生命延续；"
     "这不会带来希望契机，却逼出复杂真实的内心写照。",
     {"size": 15, "color": TEXT, "space_after": 16}),
    ("答题一句", {"size": 15, "bold": True, "color": ACCENT2, "space_after": 8}),
    ("她既摆脱不了世俗碰撞与束缚，又渴望在被扼杀的本性中保留追求幸福的自由——"
     "二者难调和，正是人物悲剧性与立体感的来源。",
     {"size": 14, "color": MUTED, "space_after": 0}),
])

# ========== 二、艺术手法总览 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "二 / 艺术手法", "写人与表现手法一览",
               "笔记提示：对比衬托（镜像）· 反讽 · 白描 · 心理 · 语言动作 · 象征")
techs = [
    ("对比／镜像", "饿嫁宣言↔下嫁\n百乐门↔夜巴黎\n朱凤↔年轻金兆丽"),
    ("反讽", "玉观音名号\n喜剧收场实为妥协\n货腰规训复现创伤"),
    ("白描", "浓脂艳粉／鱼尾纹\n假屁股假奶\n动作狠准"),
    ("心理／意识流", "镜前回忆\n由朱凤想到堕胎\n由少年想到月如"),
    ("语言描写", "娘个冬采\n尖酸骂战\n粗口＝硬壳"),
    ("象征", "钻戒／镜子\n舞厅空间\n一二三拍子"),
]
for i, (t, b) in enumerate(techs):
    col, row = i % 3, i // 3
    x = Inches(0.45) + Inches(col * 4.25)
    y = Inches(1.95) + Inches(row * 2.5)
    add_card(slide, x, y, Inches(4.05), Inches(2.3))
    add_textbox(slide, x + Inches(0.2), y + Inches(0.25), Inches(3.6), Inches(0.4),
                t, size=16, bold=True, color=ACCENT)
    add_paras(slide, x + Inches(0.2), y + Inches(0.8), Inches(3.6), Inches(1.3), [
        (line, {"size": 13, "color": MUTED, "space_after": 4})
        for line in b.split("\n")
    ])

# ========== 对比镜像 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "二 / 艺术手法", "对比与镜像：人物互相映照",
               "任黛黛／朱凤／筱红美／吴喜奎——都有金兆丽的影子")
rows = [
    ("朱凤", "怀孕哭诉、要生下孩子", "镜像年轻金兆丽：为爱奋不顾身，即将重演创伤"),
    ("任黛黛", "早嫁棉纱大王、摇扇发福", "曾被嘲，今成羡慕对象——道路选择的反面教材／样本"),
    ("筱红美", "心黑手辣、不栽跟头", "功利主义完成态；金教人不动真情的活标本"),
    ("吴喜奎", "昔日姐妹，后遁入佛堂", "另一条出路：逃避欢场，仍难逃精神荒原"),
]
for i, (name, ev, fx) in enumerate(rows):
    y = Inches(1.9) + Inches(i * 1.25)
    add_card(slide, Inches(0.5), y, Inches(12.35), Inches(1.1))
    add_textbox(slide, Inches(0.75), y + Inches(0.15), Inches(1.8), Inches(0.35),
                name, size=16, bold=True, color=ACCENT)
    add_textbox(slide, Inches(2.7), y + Inches(0.15), Inches(9.8), Inches(0.35),
                ev, size=13, color=QUOTE)
    add_textbox(slide, Inches(2.7), y + Inches(0.55), Inches(9.8), Inches(0.4),
                fx, size=13, color=MUTED)

# ========== 反讽与语言 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "二 / 艺术手法", "反讽 · 语言 · 动作白描",
               "喜剧外壳包裹悲剧内核——粗口与狠动作是生存硬壳")
quote_block(slide, Inches(0.5), Inches(1.95), Inches(6.05), Inches(2.3),
            "玩是玩，耍是耍，货腰娘第一大忌是让人家睡大肚皮。",
            "反讽：规训后辈＝复现自我创伤；真情被命名为毒药")
quote_block(slide, Inches(6.8), Inches(1.95), Inches(6.05), Inches(2.3),
            "娘个冬采！……把手包豁啷摔到化妆台上，一屁股坐在大化妆镜前，狠狠啐了一口。",
            "语言＋动作：泼辣声口瞬间立住人物；喜剧感与沧桑并出")
quote_block(slide, Inches(0.5), Inches(4.5), Inches(12.35), Inches(2.35),
            "眼角子上突然便现出了几把鱼尾巴来……勒肚子束腰，假屁股假奶。",
            "白描效果：衰老与身体改造写得具体刺目；「玉观音」神话被镜子拆穿")

# ========== 象征 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "二 / 艺术手法", "象征手法：物象承载意义",
               "钻戒／镜子／金饰／舞步——可见之物指向不可见之心")
syms = [
    ("一克拉半钻戒", "物质语言完成的补偿与救赎", "骂后甩戒：温情仍用金钱表达；不能自救却救别人"),
    ("化妆镜／鱼尾纹", "身份认同的撕裂", "镜映衰老现实；浓妆＝职业面具"),
    ("金碧辉煌首饰", "物欲追逐与舞台身份", "耳坠项链手串发针：把「金」字写在身上"),
    ("「一二三」拍子", "告别仪式／倒计时反思", "与酷似月如者共舞：卸下面具的瞬间，结局开放"),
]
for i, (t, s, n) in enumerate(syms):
    y = Inches(1.9) + Inches(i * 1.25)
    add_card(slide, Inches(0.5), y, Inches(12.35), Inches(1.1))
    add_textbox(slide, Inches(0.75), y + Inches(0.15), Inches(3.2), Inches(0.35),
                t, size=15, bold=True, color=ACCENT)
    add_textbox(slide, Inches(4.1), y + Inches(0.15), Inches(8.4), Inches(0.35),
                s, size=13, color=QUOTE)
    add_textbox(slide, Inches(0.75), y + Inches(0.55), Inches(11.8), Inches(0.4),
                n, size=12, color=MUTED)

# ========== 三、结构 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "三 / 小说结构", "结构总览：一夜框住二十年",
               "倒叙起笔＋插叙转折＋意识流缝合＋开放收束")
structs = [
    ("时间框", "现实：最后一夜\n心理：二十年欢场\n浓缩结账"),
    ("倒叙", "从西门町华灯\n「最后一夜」切入\n再回看前半生"),
    ("插叙三点", "月如爱情／堕胎\n秦雄真心等不起\n决定嫁陈发荣"),
    ("开放结局", "与少年共舞\n曲终人散留白\n孤寂余韵"),
]
for i, (t, b) in enumerate(structs):
    x = Inches(0.45) + Inches(i * 3.2)
    add_card(slide, x, Inches(2.0), Inches(3.05), Inches(4.7))
    add_textbox(slide, x + Inches(0.2), Inches(2.25), Inches(2.65), Inches(0.45),
                t, size=17, bold=True, color=ACCENT)
    add_paras(slide, x + Inches(0.2), Inches(2.9), Inches(2.65), Inches(3.5), [
        (line, {"size": 13, "color": MUTED, "space_after": 6})
        for line in b.split("\n")
    ])

# ========== 倒叙插叙 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "三 / 小说结构", "倒叙与插叙：为何从「最后一夜」写起？",
               "不直接宣泄二十年苦水，而让意识牵出关键转折，解释她如何变成今日金大班")
add_card(slide, Inches(0.5), Inches(1.95), Inches(12.35), Inches(4.9))
add_paras(slide, Inches(0.85), Inches(2.2), Inches(11.7), Inches(4.4), [
    ("三个插叙转折点（笔记）", {"size": 16, "bold": True, "color": ACCENT, "space_after": 10}),
    ("1. 盛月如：完整纯洁的爱情体验；身份悬殊；堕胎代价——情根与创伤源头。",
     {"size": 14, "color": TEXT, "space_after": 8}),
    ("2. 秦雄：真心实意，却无殷实经济；她已不能为爱再等五年。",
     {"size": 14, "color": TEXT, "space_after": 8}),
    ("3. 陈发荣：六十大几土财主；多走二十年远路后，明白「需要」压过「想要」。",
     {"size": 14, "color": TEXT, "space_after": 14}),
    ("结构效果：当下叙事与过往回忆交织，形成今昔对照与镜像；读者理解其不舍与无奈，产生共情。",
     {"size": 14, "color": MUTED, "space_after": 0}),
])

# ========== 意识流与视角 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "三 / 小说结构", "意识流与叙述视角",
               "情节随意识流动展开；全知外视点 ↔ 个人内意识自由切换")
triggers = [
    ("化妆镜鱼尾纹", "→ 失去的青春"),
    ("即将嫁陈发荣", "→ 等不起的秦雄"),
    ("怀孕的朱凤", "→ 为月如堕胎的自己"),
    ("红舞女筱红美", "→ 昔日百乐门红火"),
    ("怯态男青年", "→ 一见倾心的月如"),
]
for i, (a, b) in enumerate(triggers):
    y = Inches(1.9) + Inches(i * 0.85)
    add_card(slide, Inches(0.5), y, Inches(12.35), Inches(0.75))
    add_textbox(slide, Inches(0.8), y + Inches(0.18), Inches(5), Inches(0.4),
                a, size=15, bold=True, color=ACCENT)
    add_textbox(slide, Inches(6.0), y + Inches(0.18), Inches(6.5), Inches(0.4),
                b, size=15, color=TEXT)

# ========== 开放结局 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "三 / 小说结构", "开放式结局与留白",
               "「一二三」拍子＝舞步节奏，亦象征对人生的倒计时式反思")
quote_block(slide, Inches(0.5), Inches(1.95), Inches(12.35), Inches(2.0),
            "结尾与酷似月如的年轻人共舞，轻轻柔柔地数着拍子「一二三」……曲终人散。",
            "引证／笔记：未明确交代未来；留下孤寂与悲凉余韵，由读者填补命运可能")
add_card(slide, Inches(0.5), Inches(4.2), Inches(12.35), Inches(2.65))
add_paras(slide, Inches(0.85), Inches(4.45), Inches(11.7), Inches(2.2), [
    ("为何这样收？", {"size": 15, "bold": True, "color": ACCENT, "space_after": 8}),
    ("• 若写明婚后幸福／痛苦，主题会被说死；留白强化悲剧美学。",
     {"size": 14, "color": TEXT, "space_after": 6}),
    ("• 共舞瞬间仿佛卸下荡妇面具，短暂碰触未被玷污的自我——然后仍须面对「棺材板」婚姻。",
     {"size": 14, "color": TEXT, "space_after": 6}),
    ("• 效果：灵肉冲突无解；喜剧外衣下的苍凉收束。",
     {"size": 14, "color": MUTED, "space_after": 0}),
])

# ========== 四、环境总览 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "四 / 环境", "环境三层：微观舞厅 · 历史背景 · 象征空间",
               "舞池小舞台浓缩大时代风云；个人今不如昔＝时代缩影")
envs = [
    ("微观环境", "百乐门盛 ↔ 夜巴黎衰\n金兆丽自身今不如昔\n客居台湾的怀旧感伤"),
    ("历史背景", "1940–60s 迁徙\n上海→台北流亡\n《台北人》时代挽歌"),
    ("象征隐喻", "舞厅＝浮华虚幻\n服饰身体＝物欲面具\n镜子／月如＝青春不可追"),
]
for i, (t, b) in enumerate(envs):
    x = Inches(0.5) + Inches(i * 4.2)
    add_card(slide, x, Inches(2.05), Inches(4.0), Inches(4.7))
    add_textbox(slide, x + Inches(0.25), Inches(2.3), Inches(3.5), Inches(0.45),
                t, size=18, bold=True, color=ACCENT)
    add_paras(slide, x + Inches(0.25), Inches(3.0), Inches(3.5), Inches(3.4), [
        (line, {"size": 14, "color": MUTED, "space_after": 8})
        for line in b.split("\n")
    ])

# ========== 微观环境引证 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "四 / 环境", "微观环境：百乐门 vs 夜巴黎",
               "三次回忆把「盛世难再」钉死——空间对比即价值落差")
quote_block(slide, Inches(0.5), Inches(1.95), Inches(12.35), Inches(1.55),
            "夜巴黎不靠了我玉观音金兆丽这块老牌子，就撑得起今天这个场面了？",
            "① 名气撑场：今日排场仍吃往昔品牌老本")
quote_block(slide, Inches(0.5), Inches(3.7), Inches(6.05), Inches(3.1),
            "百乐门里那间厕所只怕比夜巴黎的舞池还宽敞些呢，童得怀那副嘴脸在百乐门掏粪坑未必有他的份。",
            "② 空间羞辱式对比：台北舞厅远不及大上海繁华")
quote_block(slide, Inches(6.8), Inches(3.7), Inches(6.05), Inches(3.1),
            "哪里见过从前那种日子？那种架势？当年在上海，拜倒她玉观音裙下，像陈发荣那点根基的人，扳起脚趾头来还数不完呢！",
            "③ 人物今昔：陈发荣类在当年不稀罕；眼浅小娼妇未见过大世面")

# ========== 历史背景 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "四 / 环境", "历史背景：迁徙与异乡人的悲苦",
               "笔记：作家是时代记录者；金兆丽命运＝大时代风流云散的微缩")
add_card(slide, Inches(0.5), Inches(1.95), Inches(12.35), Inches(4.9))
add_paras(slide, Inches(0.85), Inches(2.2), Inches(11.7), Inches(4.4), [
    ("背景要点", {"size": 16, "bold": True, "color": ACCENT, "space_after": 10}),
    ("• 时空：20世纪40–60年代；地点从命运多舛的上海到客居的台北。",
     {"size": 14, "color": TEXT, "space_after": 8}),
    ("• 人群：败退下的大陆客——将军夫人到风尘女子，皆有「大陆辉煌／台湾不如意」。",
     {"size": 14, "color": TEXT, "space_after": 8}),
    ("• 情感轨迹：怀希望 → 经历失望 → 近于绝望；在美好回忆与现世挣扎间徘徊。",
     {"size": 14, "color": TEXT, "space_after": 8}),
    ("• 系列主题接口：没落与时代挽歌；对异乡人无处安放之灵魂的同情。",
     {"size": 14, "color": TEXT, "space_after": 14}),
    ("效果：以小见大——一个货腰娘的收场夜，映出一代人繁华落尽的悲肠。",
     {"size": 14, "color": MUTED, "space_after": 0}),
])

# ========== 象征环境 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "四 / 环境", "象征与隐喻：环境如何「说话」",
               "舞厅不是中性布景，而是浮华、商品化与逃避的符号空间")
add_card(slide, Inches(0.5), Inches(1.95), Inches(4.0), Inches(4.9))
add_paras(slide, Inches(0.75), Inches(2.2), Inches(3.5), Inches(4.4), [
    ("舞厅空间", {"size": 16, "bold": True, "color": ACCENT, "space_after": 10}),
    ("灯光／音乐／舞步＝梦幻氛围", {"size": 13, "color": TEXT, "space_after": 6}),
    ("掩盖孤独与无奈", {"size": 13, "color": MUTED, "space_after": 10}),
    ("客人寻短暂欢愉＝逃避现实", {"size": 13, "color": TEXT, "space_after": 6}),
    ("朱凤揭示：舞女是商品也是牺牲品", {"size": 13, "color": MUTED, "space_after": 0}),
])
add_card(slide, Inches(4.7), Inches(1.95), Inches(4.0), Inches(4.9))
add_paras(slide, Inches(4.95), Inches(2.2), Inches(3.5), Inches(4.4), [
    ("服饰与身体", {"size": 16, "bold": True, "color": ACCENT, "space_after": 10}),
    ("金碧辉煌装扮＝物欲", {"size": 13, "color": TEXT, "space_after": 6}),
    ("大红大紫 → 结局淡化", {"size": 13, "color": MUTED, "space_after": 10}),
    ("共舞时轻轻数拍", {"size": 13, "color": TEXT, "space_after": 6}),
    ("隐喻：卸下面具，短暂回归本真", {"size": 13, "color": MUTED, "space_after": 0}),
])
add_card(slide, Inches(8.9), Inches(1.95), Inches(4.0), Inches(4.9))
add_paras(slide, Inches(9.15), Inches(2.2), Inches(3.5), Inches(4.4), [
    ("镜子与月如", {"size": 16, "bold": True, "color": ACCENT, "space_after": 10}),
    ("镜中鱼尾纹＝衰老现实", {"size": 13, "color": TEXT, "space_after": 6}),
    ("身份撕裂的视觉证据", {"size": 13, "color": MUTED, "space_after": 10}),
    ("月如年轻身体＝不可追回的青春符号", {"size": 13, "color": TEXT, "space_after": 6}),
    ("双重反射：情与时皆不可返", {"size": 13, "color": MUTED, "space_after": 0}),
])

# ========== 综合收束 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
section_header(slide, "综合", "证据 → 手法／结构／环境 → 人物与主题",
               "笔记主旨接口：今昔之变 · 灵肉之争 · 物质与精神 · 时代挽歌")
add_card(slide, Inches(0.5), Inches(1.95), Inches(12.35), Inches(4.9))
add_paras(slide, Inches(0.85), Inches(2.2), Inches(11.7), Inches(4.4), [
    ("四块拼图如何合拢", {"size": 16, "bold": True, "color": ACCENT, "space_after": 12}),
    ("人物：矛盾立体的金兆丽（道路／处世／爱情三线拧在一起）。",
     {"size": 14, "color": TEXT, "space_after": 8}),
    ("手法：对比镜像、反讽、粗口白描、意识流、象征物象把矛盾「做出来」。",
     {"size": 14, "color": TEXT, "space_after": 8}),
    ("结构：一夜框一生；倒叙插叙解释她如何走到下嫁前夜；开放结局留下苍凉。",
     {"size": 14, "color": TEXT, "space_after": 8}),
    ("环境：夜巴黎对百乐门的坠落，把个人收场嵌进迁徙时代的挽歌。",
     {"size": 14, "color": TEXT, "space_after": 14}),
    ("总括：喜剧外壳下，是一个货腰娘在繁华落尽后的屈降、残存真情，与一代异乡人的感伤。",
     {"size": 14, "color": MUTED, "space_after": 0}),
])

# ========== 尾页 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_bar(slide, 0, 0, W, Inches(0.12), ACCENT)
add_textbox(slide, Inches(0.8), Inches(2.6), Inches(11.5), Inches(0.8),
            "《金大班的最后一夜》", size=36, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_textbox(slide, Inches(0.8), Inches(3.5), Inches(11.5), Inches(0.45),
            "人物 · 手法 · 结构 · 环境", size=20, color=ACCENT, align=PP_ALIGN.CENTER)
add_textbox(slide, Inches(0.8), Inches(4.3), Inches(11.5), Inches(0.4),
            "精读须有细节引证 —— 声口、镜子、钻戒、百乐门与夜巴黎",
            size=14, color=MUTED, align=PP_ALIGN.CENTER)

out1 = "/workspace/金大班的最后一夜_精读教学课件.pptx"
out2 = "/workspace/JinDaban_LastNight_PPT.pptx"
prs.save(out1)
import shutil
shutil.copy(out1, out2)
shutil.copy(out1, "/opt/cursor/artifacts/JinDaban_LastNight_PPT.pptx")
shutil.copy(out1, "/opt/cursor/artifacts/金大班的最后一夜_精读教学课件.pptx")
print(f"OK slides={len(prs.slides)}")
print(out1)
